"""Enumeration used to be silent about what it could not reach.

That silence is the whole problem. A pack missing the ending nobody could get to looks exactly
like a pack that is complete, so the reading is a suspicion -- "some terminal states are not
reached within a capability" -- with no way to tell whether the declaration is wrong, the span is
drawn wrong, or the walk is. Three different fixes, and no way to choose between them.

Naming the ending and the reason turns that into something to act on. Answered by reachability
rather than by enumerating, so it costs two searches per block and can be shown on a page that
redraws rather than only after a run.
"""
import dataclasses
import unittest
from pathlib import Path

from scenario_generator.core import graph as graph_module
from scenario_generator.core.graph import DecisionGraph, endings_not_reached
from scenario_generator.core.intake import read_intake

EXAMPLE = (Path(__file__).resolve().parent.parent
           / "examples" / "intakes" / "2_travel_no_spans.xlsx")


class TestNothingIsReportedWhereNothingIsWrong(unittest.TestCase):
    """The first thing a warning has to be is quiet when there is nothing to warn about."""

    def test_a_well_drawn_declaration_reports_nothing(self):
        for name in ("1_disputes_three_blocks", "2_travel_no_spans"):
            path = EXAMPLE.parent / f"{name}.xlsx"
            intake = read_intake(str(path))
            graph = DecisionGraph(intake.decisions, intake.states)
            self.assertEqual(endings_not_reached(graph, intake.capabilities, intake.decisions), [],
                             name)


class TestTheThreeReasons(unittest.TestCase):
    def setUp(self):
        self.intake = read_intake(str(EXAMPLE))
        self.graph = DecisionGraph(self.intake.decisions, self.intake.states)

    def _spanned(self, capability_id, entry, exits):
        return tuple(dataclasses.replace(c, entry_states=(entry,), exit_states=tuple(exits))
                     if c.id == capability_id else dataclasses.replace(
                         c, entry_states=(), exit_states=())
                     for c in self.intake.capabilities)

    def _found(self, capabilities):
        return {u.state_id: u.reason
                for u in endings_not_reached(self.graph, capabilities, self.intake.decisions)}

    def test_an_exit_drawn_across_the_middle_of_a_block(self):
        """S-01 sits between DEC-01 and DEC-02, so a block told to hand on there never reaches
        what DEC-02 ends at. The span is too narrow -- the declaration is fine."""
        found = self._found(self._spanned("CAP-01", "S-00", ["S-01"]))
        self.assertIn("S-04", found)
        self.assertIn("hand on at S-01", found["S-04"])

    def test_a_block_entered_past_its_own_decisions(self):
        """Entered at S-03, CAP-01's decisions are behind it, so its endings are unreachable.
        This is the one that means the intake or the entry is wrong."""
        found = self._found(self._spanned("CAP-01", "S-03", ["S-07"]))
        self.assertIn("S-02", found)
        self.assertIn("cannot be reached", found["S-02"])

    def test_a_route_longer_than_the_backstop(self):
        """Nothing wrong with the declaration; the enumeration is incomplete and says so."""
        original = graph_module.MAX_DEPTH
        graph_module.MAX_DEPTH = 1
        try:
            found = self._found(self._spanned("CAP-01", "S-00", ["S-03"]))
        finally:
            graph_module.MAX_DEPTH = original
        self.assertIn("S-04", found)
        self.assertIn("past the 1-step backstop", found["S-04"])

    def test_it_names_the_capability(self):
        found = endings_not_reached(self.graph, self._spanned("CAP-01", "S-00", ["S-01"]),
                                    self.intake.decisions)
        self.assertEqual({u.capability_id for u in found}, {"CAP-01"})


class TestRetryBoundsAreNotReported:
    """Placeholder for the reasoning, kept out of the run deliberately.

    An ending only reachable past a retry bound *is* reached -- by the augmentation pass, which
    exists for exactly that -- so reporting it here would name a route that does get walked.
    Reachability ignores attempt counts for that reason, which is the same choice
    :func:`core.graph._decisions_within` makes and for the same reason.
    """


class TestItIsSaidBesideTheDrawing(unittest.TestCase):
    """Beside the graph rather than only after a run, because the fix for every one of these
    reasons is an edit to the declaration and the declaration is on that page."""

    def _page(self, capabilities):
        import shutil
        import tempfile
        import time

        from openpyxl import load_workbook

        from scenario_generator.webapp.app import create_app

        root = Path(tempfile.mkdtemp())
        source = root / "intake.xlsx"
        shutil.copy(EXAMPLE, source)
        book = load_workbook(source)
        sheet = book["L2 Capabilities"]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == "CAP-01":
                row[3].value, row[4].value = capabilities
        book.save(source)

        client = create_app(root / "ws").test_client()
        client.post("/workspaces", data={"name": "Reach"})
        with open(source, "rb") as handle:
            client.post("/stage/intake/upload",
                        data={"files": (handle, "intake.xlsx"), "group": "intake_workbook"},
                        content_type="multipart/form-data")
        client.post("/stage/intake/run")
        for _ in range(600):
            if client.get("/stage/intake/progress").get_json()["status"] != "running":
                break
            time.sleep(0.05)
        return client.get("/stage/intake").get_data(as_text=True)

    def test_the_ending_and_the_reason_are_both_on_the_page(self):
        page = self._page(("S-00", "S-01"))
        self.assertIn("No route ends here", page)
        self.assertIn("S-04", page)
        self.assertIn("hand on at", page)

    def test_a_sound_declaration_says_nothing(self):
        self.assertNotIn("No route ends here", self._page(("S-00", "S-03")))


class TestTheBackstopFitsARealDeclaration(unittest.TestCase):
    def test_depth_is_not_set_under_a_long_route(self):
        """It was 12. Thirty-odd states with a three-attempt retry in the middle spends steps
        quickly, and every branch past the twelfth was dropped with nothing on screen to say so."""
        self.assertGreaterEqual(graph_module.MAX_DEPTH, 24)


if __name__ == "__main__":
    unittest.main()
