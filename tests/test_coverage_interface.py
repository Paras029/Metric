"""The coverage stage as a person actually uses it: the counts on the page, the threshold, the pack.

Three properties are pinned here, all of them ones that would fail quietly.

The threshold is a *view* of the mappings, not a rerun. Moving it must recount what is already
stored and rewrite the workbook without another model call, because a person arrives at the right
number by trying one and looking at the result -- and nobody tries a second number if the first
costs a pass over every transcript.

The counts and the confidences stay separate all the way to the page. Folding a weak match into
the figure would decide, behind the reader, something they are far better placed to decide.

The pack filter subtracts, which nothing else in this pipeline does. It must be off by default,
must never narrow the metadata, and must not empty the pack when coverage has not run.
"""
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from metric.web.server import create_app
from metric.web.workspace import Workspace

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_stage_runners import _intake_workbook  # noqa: E402  the same minimal valid intake


def _conversations(directory: Path, count: int) -> Path:
    path = directory / "their_conversations.xlsx"
    book = Workbook()
    book.active.append(["Conversation ID", "Turn", "Speaker", "Message", "Scenario"])
    for number in range(1, count + 1):
        book.active.append([f"C{number}", 1, "Customer", "I need into my account", "Sign in"])
        book.active.append([f"C{number}", 2, "Agent", "You are authenticated now", "Sign in"])
    book.save(path)
    return path


class TestTheCoveragePage(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.app = create_app(self.root)
        self.client = self.app.test_client()
        self.client.post("/workspaces", data={"name": "Coverage check"})

        scratch = Path(tempfile.mkdtemp())
        for key, path, group in (("intake", _intake_workbook(scratch), "intake_workbook"),
                                 ("coverage", _conversations(scratch, 2), "owner_scenarios")):
            with open(path, "rb") as handle:
                self.client.post(f"/stage/{key}/upload",
                                 data={"files": (handle, path.name), "group": group},
                                 content_type="multipart/form-data")
        self._run("intake")
        self._run("workflow")

    def _run(self, key):
        self.client.post(f"/stage/{key}/run")
        for _ in range(600):
            if self.client.get(f"/stage/{key}/progress").get_json()["status"] != "running":
                return
            import time
            time.sleep(0.05)
        self.fail(f"{key} never finished")

    def _workspace(self):
        return Workspace.load(self.root / "coverage-check")

    def _first_scenario(self):
        from metric.domain import read_space_metadata
        return read_space_metadata(str(self._workspace().root / "scenario_space_metadata.xlsx"))[0].id

    def _map(self, calls=None):
        """Run coverage with both conversations landing on the first scenario."""
        target = self._first_scenario()
        reply = json.dumps({f"C{n}": {"scenario_id": target, "confidence": "high" if n == 1
                                      else "low", "reason": "ends signed in"}
                            for n in (1, 2)})

        def stub(system, user, **kwargs):
            if calls is not None:
                calls.append(user)
            return reply

        with mock.patch("metric.phases.coverage.coverage.mapping.ask_llm", stub):
            self._run("coverage")
        return target

    def test_the_page_shows_the_count_and_the_confidence_split_side_by_side(self):
        target = self._map()
        page = self.client.get("/stage/coverage").data.decode()
        self.assertIn("2 conversations", page)
        self.assertIn("1 high, 1 low", page)
        self.assertIn(target, page)

    def test_a_scenario_nothing_reached_is_shown_rather_than_omitted(self):
        self._map()
        page = self.client.get("/stage/coverage").data.decode()
        self.assertIn("Never exercised", page)

    def test_moving_the_threshold_recounts_without_calling_the_model_again(self):
        self._map()
        before = self._workspace().coverage_mappings()

        calls = []
        with mock.patch("metric.phases.coverage.coverage.mapping.ask_llm",
                        lambda *a, **k: calls.append(a) or "{}"):
            self.client.post("/stage/coverage/threshold", data={"threshold": "3"})

        self.assertEqual(calls, [])
        self.assertEqual(self._workspace().coverage_threshold, 3)
        self.assertEqual(self._workspace().coverage_mappings(), before)

    def test_the_threshold_changes_what_counts_as_represented(self):
        """Two conversations on one scenario: represented at two, under-represented at three."""
        self._map()
        self.assertEqual(self._under_represented_after(1), 1)   # only the untouched scenario
        self.assertEqual(self._under_represented_after(2), 1)
        self.assertEqual(self._under_represented_after(3), 2)   # the covered one drops below

    def _under_represented_after(self, threshold):
        """Set the threshold through the route, then count what the page would now show."""
        from metric.domain import read_space_metadata
        from metric.phases.coverage.coverage.view import stored_report

        self.client.post("/stage/coverage/threshold", data={"threshold": str(threshold)})
        workspace = self._workspace()
        report = stored_report(workspace, read_space_metadata(str(workspace.root / "scenario_space_metadata.xlsx")))
        return len(report.under_represented())

    def test_the_workbook_is_rewritten_at_the_new_threshold(self):
        self._map()
        self.client.post("/stage/coverage/threshold", data={"threshold": "5"})

        book = load_workbook(self._workspace().root / "coverage.xlsx")
        rows = list(book["Scenarios"].iter_rows(min_row=1, values_only=True))
        # Read by header rather than by position. A column added to this sheet is a change to the
        # report, not to what "represented" means, and a test that reads column five turns the
        # first into the second.
        at = {str(name): index for index, name in enumerate(rows[0])}
        verdicts = {row[at["SC ID"]]: row[at["Represented?"]] for row in rows[1:]}
        self.assertTrue(all(v == "No" for v in verdicts.values()),
                        "nothing has five conversations, so nothing is represented")

    def test_the_registry_records_the_count_against_the_scenario(self):
        target = self._map()
        book = load_workbook(self._workspace().root / "scenario_space_metadata.xlsx")
        sheet = book["Scenario_Metadata"]
        headers = [c.value for c in sheet[1]]
        column = headers.index("Owner Coverage")
        rows = {row[0]: row[column] for row in sheet.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows[target], "2 conversations")


class TestWhatGoesInThePack(unittest.TestCase):
    """The one control in the pipeline that removes scenarios rather than adding them."""

    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.client = create_app(self.root).test_client()
        self.client.post("/workspaces", data={"name": "Pack scope"})
        scratch = Path(tempfile.mkdtemp())
        with open(_intake_workbook(scratch), "rb") as handle:
            self.client.post("/stage/intake/upload",
                             data={"files": (handle, "intake.xlsx"), "group": "intake_workbook"},
                             content_type="multipart/form-data")
        for key in ("intake", "workflow"):
            self.client.post(f"/stage/{key}/run")
            self._settle(key)

    def _settle(self, key):
        import time
        for _ in range(600):
            if self.client.get(f"/stage/{key}/progress").get_json()["status"] != "running":
                return
            time.sleep(0.05)
        self.fail(f"{key} never finished")

    def _workspace(self):
        return Workspace.load(self.root / "pack-scope")

    def _issue(self):
        self.client.post("/stage/summary/run")
        self._settle("summary")
        book = load_workbook(self._workspace().root / "data_template.xlsx")
        return book["Scenarios"].max_row - 1

    def test_the_whole_space_is_issued_by_default(self):
        self.assertFalse(self._workspace().pack_gaps_only)
        self.assertGreater(self._issue(), 0)

    def test_turning_it_on_with_no_coverage_run_still_issues_everything(self):
        """Nothing mapped is not the same as everything covered, and emptying the pack on the
        strength of a stage that never ran would be the worst outcome available."""
        everything = self._issue()
        self.client.post("/stage/summary/scope", data={"on": "1"})
        self.assertTrue(self._workspace().pack_gaps_only)
        self.assertEqual(self._issue(), everything)

    def test_a_covered_scenario_is_held_back_but_stays_in_the_registry(self):
        from metric.domain import read_space_metadata
        workspace = self._workspace()
        metadata = read_space_metadata(str(workspace.root / "scenario_space_metadata.xlsx"))
        covered = metadata[0].id

        workspace.save_coverage([_mapping("C1", covered), _mapping("C2", covered)], "test")
        workspace.pack_gaps_only = True
        workspace.save()

        issued = self._issue()
        self.assertEqual(issued, len(metadata) - 1)
        self.assertEqual(len(read_space_metadata(str(self._workspace().root / "scenario_space_metadata.xlsx"))),
                         len(metadata))


def _mapping(conversation_id: str, scenario_id: str):
    from metric.phases.coverage.coverage.mapping import Mapping
    return Mapping(conversation_id=conversation_id, scenario_id=scenario_id, confidence="high")


if __name__ == "__main__":
    unittest.main()
