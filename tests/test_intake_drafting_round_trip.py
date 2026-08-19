"""What the drafter extracts has to survive into the workbook, and be checked once it is there.

The intake is the boundary of what can be tested, so a field the drafter read and the writer lost
is a part of the agent nothing downstream ever hears about -- and it is lost silently, which is the
property that makes it worth its own test file. Two guards.

*Everything extracted is written.* The workbook template carries fields the drafter is not asked
for, so the two lists are different lengths and in different orders. Anything that walks them in
step is wrong, and wrong in the quiet direction: the cell it should have written stays empty.

*What the draft leaves structurally broken is put back to the model.* One call over a long context
routinely leaves a branch with a single named outcome or an outcome leading to a state nobody
declared. Neither is a matter of opinion -- the graph cannot be walked without them -- and the
answer is usually a paragraph away in the documentation the first pass had already read.
"""
import json
import tempfile
import unittest
from pathlib import Path

from scenario_generator.core import read_intake
from scenario_generator.core.intake import write_template
from scenario_generator.ingest import DraftedIntake, write_drafted_intake
from scenario_generator.ingest.drafting import _L1_KEYS, _validate
from scenario_generator.pipeline import draft_intake_workbook, structural_problems

_USE_CASE = {
    "name": "Claims Assistant",
    "objective": "File motor claims without a person for the routine cases",
    "agent_type": "Chatbot in the servicing app",
    "channel": "Mobile app",
    "handoff_triggers": "Anything over £10,000, or a complaint",
    "safety_requirements": "Never quote a settlement figure",
    "success_criteria": "The claim is filed and the reference number is given",
}

_MINIMAL = {"use_case": _USE_CASE, "personas": [], "capabilities": [], "decisions": [],
            "states": [], "tools": []}


def _written(data: dict) -> dict:
    """One draft written to a workbook and read back, as the pipeline would."""
    path = Path(tempfile.mkdtemp()) / "intake.xlsx"
    write_drafted_intake(path, DraftedIntake(_validate(data)))
    return read_intake(str(path)).use_case


class TestNothingExtractedIsLostOnTheWayToTheWorkbook(unittest.TestCase):
    def test_every_use_case_field_the_drafter_returns_reaches_the_workbook(self):
        use_case = _written(_MINIMAL)
        self.assertEqual(use_case["Use case name"], _USE_CASE["name"])
        self.assertEqual(use_case["Business objective"], _USE_CASE["objective"])
        self.assertEqual(use_case["Agent type"], _USE_CASE["agent_type"])
        self.assertEqual(use_case["Channel / modality"], _USE_CASE["channel"])
        self.assertEqual(use_case["Human handoff triggers"], _USE_CASE["handoff_triggers"])
        self.assertEqual(use_case["Safety requirements"], _USE_CASE["safety_requirements"])
        self.assertEqual(use_case["Success criteria"], _USE_CASE["success_criteria"])

    def test_the_labels_the_writer_looks_for_are_labels_the_template_has(self):
        """The writer finds each field by its label. A label the template does not carry is a
        field that can never be written, and nothing else in the round trip would say so."""
        from openpyxl import load_workbook

        path = Path(tempfile.mkdtemp()) / "template.xlsx"
        write_template(str(path))
        sheet = load_workbook(path)["L1 Use Case"]
        labels = {str(sheet.cell(row=row, column=1).value or "").strip()
                  for row in range(2, sheet.max_row + 1)}
        for label, _key in _L1_KEYS:
            self.assertIn(label, labels, f"the template has no '{label}' row to write into")

    def test_a_field_the_drafter_left_out_is_blank_rather_than_shifting_the_rest(self):
        """The failure mode worth ruling out is not a blank cell -- it is a blank cell that also
        moved every field after it into the wrong row."""
        partial = dict(_MINIMAL, use_case={"name": "Claims Assistant",
                                           "success_criteria": "The claim is filed"})
        use_case = _written(partial)
        self.assertEqual(use_case["Use case name"], "Claims Assistant")
        self.assertEqual(use_case["Success criteria"], "The claim is filed")
        self.assertEqual(use_case["Agent type"], "")
        self.assertEqual(use_case["Business objective"], "")


# A declaration that reads back cleanly and cannot be walked: one decision with a single named
# outcome, leading to a state nothing declares.
_BROKEN = {
    "use_case": _USE_CASE,
    "personas": [{"id": "P1", "name": "Cooperative user",
                  "applies_to": "Wants to file a claim and is trying to", "is_default": True}],
    "capabilities": [{"id": "CAP-01", "name": "Policy check", "type": "Gating"}],
    "decisions": [{"id": "DEC-01", "name": "Policy in force?", "capability_id": "CAP-01",
                   "inputs": "policy number", "outcomes": ["In force"], "input_source": "User",
                   "max_attempts": 1, "outcome_condition": "status is active"}],
    "states": [{"id": "S-00", "reached_via": "Start", "description": "Session begins",
                "next_decisions": ["DEC-01"], "is_terminal": False, "outcome_type": ""}],
    "tools": [],
    "review_notes": [],
}

_REPAIRED = json.loads(json.dumps(_BROKEN))
_REPAIRED["decisions"][0]["outcomes"] = ["In force", "Lapsed"]
_REPAIRED["states"] += [
    {"id": "S-01", "reached_via": "DEC-01=In force", "description": "Policy confirmed",
     "next_decisions": [], "is_terminal": True, "outcome_type": "Happy path"},
    {"id": "S-02", "reached_via": "DEC-01=Lapsed", "description": "Cannot file on a lapsed policy",
     "next_decisions": [], "is_terminal": True, "outcome_type": "Termination"}]
_REPAIRED["review_notes"] = [{"field": "DEC-01", "note": "Second outcome inferred."}]

# A complete, walkable declaration -- what a revision is given and must carry forward.
_FULL = json.loads(json.dumps(_REPAIRED))
_FULL["decisions"].append(
    {"id": "DEC-02", "name": "Cover applies?", "capability_id": "CAP-01", "inputs": "schedule",
     "outcomes": ["Covered", "Not covered"], "input_source": "Tool", "max_attempts": 1,
     "outcome_condition": "the schedule lists it"})
_FULL["states"][1]["next_decisions"] = ["DEC-02"]
_FULL["states"][1]["is_terminal"] = False
_FULL["states"][1]["outcome_type"] = ""
_FULL["states"] += [
    {"id": "S-03", "reached_via": "DEC-02=Covered", "description": "Claim opened",
     "next_decisions": [], "is_terminal": True, "outcome_type": "Happy path"},
    {"id": "S-04", "reached_via": "DEC-02=Not covered", "description": "Declined, not covered",
     "next_decisions": [], "is_terminal": True, "outcome_type": "Termination"}]


class TestTheDraftIsCheckedAndPutBack(unittest.TestCase):
    def setUp(self):
        self.work = Path(tempfile.mkdtemp())
        self.context = self.work / "run_context.md"
        self.context.write_text("# Claims\n\nA lapsed policy cannot be claimed on.\n",
                                encoding="utf-8")
        self.output = self.work / "intake.xlsx"
        self.prompts = []

    # ``reconcile=False`` throughout this class. It exercises the repair round specifically, and
    # a third call answered by the same stub would put the broken draft back over the repair.

    def _complete(self, repaired=None):
        def complete(system, user, **kwargs):
            self.prompts.append(user)
            if "WHAT IS MISSING" in user:
                return json.dumps(repaired if repaired is not None else _REPAIRED)
            return json.dumps(_BROKEN)
        return complete

    def test_a_broken_draft_is_put_back_to_the_model_with_its_own_failures_named(self):
        draft_intake_workbook(str(self.context), str(self.output), complete=self._complete(),
                              reconcile=False)

        self.assertEqual(len(self.prompts), 2, "the draft was not checked and put back")
        repair = self.prompts[1]
        self.assertIn("WHAT IS MISSING", repair)
        self.assertIn("names only one outcome", repair)
        self.assertIn("No state in the intake declares itself reached by it", repair)
        # The documents go with it: the answer is usually already in them.
        self.assertIn("A lapsed policy cannot be claimed on", repair)

    def test_the_repaired_declaration_is_what_lands_on_disk(self):
        draft_intake_workbook(str(self.context), str(self.output), complete=self._complete(),
                              reconcile=False)

        intake = read_intake(str(self.output))
        self.assertEqual(intake.decisions[0].variants, ["In force", "Lapsed"])
        self.assertEqual({s.id for s in intake.states}, {"S-00", "S-01", "S-02"})

    def test_it_leaves_the_declaration_better_than_it_found_it(self):
        draft_intake_workbook(str(self.context), str(self.output), complete=self._complete(),
                              reconcile=False)
        after = len(structural_problems(read_intake(str(self.output))))

        self.prompts.clear()
        untouched = self.work / "untouched.xlsx"
        draft_intake_workbook(str(self.context), str(untouched), complete=self._complete(),
                              repair=False)
        before = len(structural_problems(read_intake(str(untouched))))

        self.assertLess(after, before)

    def test_a_repair_that_comes_back_empty_leaves_the_first_draft_alone(self):
        """A second look may improve a declaration and must never damage one."""
        empty = {"use_case": _USE_CASE, "personas": [], "capabilities": [], "decisions": [],
                 "states": [], "tools": []}
        draft_intake_workbook(str(self.context), str(self.output), reconcile=False,
                              complete=self._complete(repaired=empty))

        intake = read_intake(str(self.output))
        self.assertEqual([d.id for d in intake.decisions], ["DEC-01"])
        self.assertEqual([s.id for s in intake.states], ["S-00"])

    def test_a_repair_that_fails_outright_leaves_the_first_draft_alone(self):
        def complete(system, user, **kwargs):
            self.prompts.append(user)
            if "WHAT IS MISSING" in user:
                raise RuntimeError("502 from the gateway")
            return json.dumps(_BROKEN)

        draft_intake_workbook(str(self.context), str(self.output), complete=complete,
                              reconcile=False)

        intake = read_intake(str(self.output))
        self.assertEqual([d.id for d in intake.decisions], ["DEC-01"])

    def test_a_complete_declaration_is_not_put_back_at_all(self):
        """A second call is worth making when there is something to fix and not otherwise."""
        def complete(system, user, **kwargs):
            self.prompts.append(user)
            return json.dumps(_REPAIRED)

        draft_intake_workbook(str(self.context), str(self.output), complete=complete,
                              reconcile=False)
        repairs = [p for p in self.prompts if "WHAT IS MISSING" in p]
        # _REPAIRED still leaves softer gaps, but the point is that a repair is only sent for
        # what the audit actually found -- never as a routine second call.
        self.assertLessEqual(len(repairs), 1)
        if repairs:
            self.assertNotIn("names only one outcome", repairs[0])


class TestARerunBuildsOnWhatIsAlreadyThere(unittest.TestCase):
    """A second run has to add to the declaration, never start it over.

    This is the difference between a re-run that helps and one that undoes the first. Everything
    that has happened since -- answers to the questions on the page, notes, a document read again
    -- is new information about a declaration that already exists. A fresh draft cannot tell a
    correction somebody made by hand from something to re-derive from nothing, so it discards it.
    """

    def setUp(self):
        from scenario_generator.webapp.workspace import Workspace

        self.workspace = Workspace.create(Path(tempfile.mkdtemp()), "Re-run")
        (self.workspace.root / "ingest_context.md").write_text(
            "# Claims\n\nA lapsed policy cannot be claimed on.\n", encoding="utf-8")
        self.calls = []

    def _record(self, kind):
        def call(*args, **kwargs):
            self.calls.append(kind)
            write_drafted_intake(Path(args[1]), DraftedIntake(_validate(_REPAIRED)))

            class _Result:
                path = args[1]
            return _Result()
        return call

    def _patched(self):
        from unittest import mock
        return (mock.patch("scenario_generator.webapp.runners.draft_intake_workbook",
                           self._record("draft")),
                mock.patch("scenario_generator.webapp.runners.revise_intake_workbook",
                           self._record("revise")))

    def test_the_first_run_drafts_and_the_second_revises(self):
        from scenario_generator.webapp.runners import _run_intake

        draft, revise = self._patched()
        with draft, revise:
            _run_intake(self.workspace)
            _run_intake(self.workspace)
            _run_intake(self.workspace)

        self.assertEqual(self.calls, ["draft", "revise", "revise"])

    def test_a_workbook_you_uploaded_is_read_and_never_written_to(self):
        from scenario_generator.webapp.runners import _run_intake

        mine = self.workspace.root / "my_intake.xlsx"
        write_drafted_intake(mine, DraftedIntake(_validate(_REPAIRED)))
        self.workspace.state("intake").artifacts["workbook"] = mine.name
        stamp = mine.stat().st_mtime_ns

        draft, revise = self._patched()
        with draft, revise:
            summary = _run_intake(self.workspace)

        self.assertEqual(self.calls, [])
        self.assertEqual(mine.stat().st_mtime_ns, stamp)
        self.assertIn("never overwritten", str(summary["This run"]))

    def test_an_upload_carrying_the_drafters_own_filename_is_still_treated_as_yours(self):
        """Otherwise "never overwrite what somebody uploaded" turns on a filename collision."""
        from scenario_generator.webapp.app import create_app

        app = create_app(self.workspace.root.parent)
        with app.test_client() as client:
            with client.session_transaction() as session:
                session["workspace"] = self.workspace.root.name
            payload = Path(tempfile.mkdtemp()) / "drafted_intake.xlsx"
            write_drafted_intake(payload, DraftedIntake(_validate(_REPAIRED)))
            client.post("/stage/intake/upload", data={
                "files": (payload.open("rb"), "drafted_intake.xlsx"),
                "group": "intake_workbook"},
                content_type="multipart/form-data")

        from scenario_generator.webapp.workspace import Workspace

        reloaded = Workspace.load(self.workspace.root)
        stored = reloaded.state("intake").artifacts["workbook"]
        self.assertNotEqual(stored, "drafted_intake.xlsx")

        from scenario_generator.webapp.runners import _run_intake

        draft, revise = self._patched()
        with draft, revise:
            summary = _run_intake(reloaded)
        self.assertEqual(self.calls, [])
        self.assertIn("never overwritten", str(summary["This run"]))

    def test_the_result_says_which_of_the_three_happened(self):
        """Three runs that look identical on the card are indistinguishable from a dead button."""
        from scenario_generator.webapp.runners import _run_intake

        draft, revise = self._patched()
        with draft, revise:
            first = _run_intake(self.workspace)
            second = _run_intake(self.workspace)

        self.assertIn("drafted", str(first["This run"]))
        self.assertIn("revised", str(second["This run"]))
        self.assertIn("Still needs an answer", first)


class TestARevisionCannotQuietlyEmptyTheDeclaration(unittest.TestCase):
    """The failure the user sees as "it forgot everything", caught before it reaches disk."""

    def setUp(self):
        from scenario_generator.pipeline import revise_intake_workbook

        self.revise = revise_intake_workbook
        self.work = Path(tempfile.mkdtemp())
        self.context = self.work / "run_context.md"
        self.context.write_text("# Claims\n\nA lapsed policy cannot be claimed on.\n",
                                encoding="utf-8")
        self.current = self.work / "intake.xlsx"
        write_drafted_intake(self.current, DraftedIntake(_validate(_FULL)))

    def test_a_revision_that_dropped_most_of_the_graph_is_discarded(self):
        collapsed = json.loads(json.dumps(_FULL))
        collapsed["decisions"] = collapsed["decisions"][:1]
        collapsed["states"] = collapsed["states"][:1]

        self.revise(str(self.current), str(self.current),
                    context_path=str(self.context),
                    complete=lambda system, user, **kwargs: json.dumps(collapsed))

        intake = read_intake(str(self.current))
        self.assertEqual(len(intake.decisions), len(_FULL["decisions"]))
        self.assertEqual(len(intake.states), len(_FULL["states"]))

    def test_a_revision_that_keeps_the_graph_is_written(self):
        added = json.loads(json.dumps(_FULL))
        added["capabilities"].append({"id": "CAP-02", "name": "Fraud check", "type": "Gating"})

        self.revise(str(self.current), str(self.current),
                    context_path=str(self.context), repair=False,
                    complete=lambda system, user, **kwargs: json.dumps(added))

        intake = read_intake(str(self.current))
        self.assertEqual({c.id for c in intake.capabilities}, {"CAP-01", "CAP-02"})

    def test_the_current_declaration_is_what_the_model_is_asked_to_revise(self):
        seen = []

        def complete(system, user, **kwargs):
            seen.append(user)
            return json.dumps(_FULL)

        self.revise(str(self.current), str(self.current), context_path=str(self.context),
                    repair=False, complete=complete)

        self.assertIn("WHAT IS CURRENTLY DECLARED", seen[0])
        self.assertIn("DEC-01", seen[0])
        self.assertIn("DEC-02", seen[0])

    def test_every_answer_reaches_the_revision(self):
        """Answers are notes, and notes are what a revision is for."""
        seen = []

        def complete(system, user, **kwargs):
            seen.append(user)
            return json.dumps(_FULL)

        self.revise(str(self.current), str(self.current), context_path=str(self.context),
                    repair=False, complete=complete,
                    notes=["(Intake) Q: What are DEC-02's outcomes? — A: Eligible / Not eligible"])

        self.assertIn("Eligible / Not eligible", seen[0])


if __name__ == "__main__":
    unittest.main()
