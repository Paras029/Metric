"""The loop that finishes a declaration, driven by stubs rather than by a gateway.

The decisions are the part worth testing, and a loop tested only against a live model is tested
on a good day. Every case here drives it with a scripted conversation, so what is being checked is
the loop's own behaviour: when it stops, what it refuses, what it does with a tool that fails.

The property the whole thing rests on: **the model never decides it is finished.** That is settled
by the audit, against the declaration on disk. A loop that judges its own completion runs until it
feels like stopping, which on a bad day is never and on a worse day is immediately.
"""
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from metric.phases.intake.intake import agent

EXAMPLES = Path(__file__).resolve().parent.parent / "examples" / "intakes"


def _scripted(*turns):
    """A conversation that returns each scripted turn once, then stops calling tools."""
    remaining = list(turns)

    def converse(messages):
        if remaining:
            return remaining.pop(0)
        return {"content": "Nothing further.", "tool_calls": [], "raw_tool_calls": []}
    return converse


def _call(_tool, **args):
    """One scripted tool call. The tool's own name is positional-only so it cannot collide with
    an argument called `name`, which two of these tools take."""
    return {"content": "", "raw_tool_calls": [],
            "tool_calls": [{"tool": _tool, "args": args, "id": f"c{_tool}"}]}


class TestWhenItStops(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.intake = self.root / "intake.xlsx"
        shutil.copy(EXAMPLES / "4_spans_drawn_wrongly.xlsx", self.intake)

    def test_a_complete_declaration_is_not_worked_on_at_all(self):
        """The cheapest correct outcome, and one a loop that trusts a model never reaches."""
        clean = self.root / "clean.xlsx"
        shutil.copy(EXAMPLES / "1_disputes_three_blocks.xlsx", clean)
        state = agent.run(self.root, str(clean), converse=_scripted())
        self.assertEqual(state.calls, 0)
        self.assertIn("already structurally complete", state.stopped_because)

    def test_the_model_saying_it_is_done_does_not_end_it(self):
        """The property everything else rests on."""
        done = {"content": "The declaration is complete.", "tool_calls": [], "raw_tool_calls": []}
        state = agent.run(self.root, str(self.intake), converse=_scripted(done, done, done))
        self.assertGreater(state.gaps_now, 0)
        self.assertNotIn("clean", state.stopped_because)

    def test_asking_for_the_same_thing_over_and_over_ends_the_sweep(self):
        """A turn that asks for exactly what the last turn asked for gets the same answer, so the
        turn after it asks again. Two of those is a loop, and a loop at judgement tier is expensive
        in a way nothing on screen makes obvious."""
        def forever(messages):
            return _call("audit_declaration")
        state = agent.run(self.root, str(self.intake), converse=forever, max_turns=24)
        self.assertLessEqual(state.turns, 4, "went round in circles at the budget's expense")

    def test_the_hard_budget_still_applies_when_the_calls_genuinely_vary(self):
        """The repetition guard is not the only stop: a model doing different things every turn
        and never finishing still has to be cut off somewhere."""
        turns = {"n": 0}

        def varied(messages):
            turns["n"] += 1
            return _call("read_document", name=f"absent-{turns['n']}.pdf")

        state = agent.run(self.root, str(self.intake), converse=varied, max_turns=5)
        self.assertEqual(state.turns, 5)
        self.assertIn("budget of 5 turns", state.stopped_because)

    def test_a_sweep_that_closes_nothing_ends_the_run(self):
        """Everything the second sweep could read, the first could read. What is left at that
        point is what the documents do not say, and another pass is money for nothing."""
        replies = [_call("what_is_declared"), _call("audit_declaration")]

        def stalls(messages):
            return replies.pop(0) if replies else {"content": "no more",
                                                   "tool_calls": [], "raw_tool_calls": []}

        state = agent.run(self.root, str(self.intake), converse=stalls, max_turns=24, sweeps=4)
        self.assertIn("closed nothing", state.stopped_because)
        self.assertLess(state.turns, 10)

    def test_a_document_is_not_sent_twice(self):
        """The text is already in the conversation; sending it again buys nothing, costs the whole
        document in tokens, and precedes repeating whatever was concluded from it."""
        sources = self.root / "sources" / "model_doc"
        sources.mkdir(parents=True)
        (sources / "notes.md").write_text("# Agent\n\nIt verifies the cardmember.\n")
        tools = {t.name: t for t in agent.build_tools(
            agent.Pack(self.root), str(self.intake), {}, {})}

        first = tools["read_document"].run(name="notes.md")
        second = tools["read_document"].run(name="notes.md")
        self.assertIn("verifies the cardmember", first)
        self.assertIn("already in this conversation", second)

    def test_the_progress_line_says_what_is_happening(self):
        """A constant line is why a working loop reads as a stuck one."""
        lines = []
        agent.run(self.root, str(self.intake),
                  converse=_scripted(_call("list_sources"), _call("audit_declaration")),
                  progress=lambda message, done, total: lines.append(message))
        self.assertIn("Looking at what was submitted", lines)
        self.assertIn("Checking what the declaration still needs", lines)
        self.assertGreater(len(set(lines)), 1, "every turn printed the same sentence")

    def test_a_failing_model_call_stops_it_without_losing_what_was_done(self):
        def breaks(messages):
            raise RuntimeError("gateway down")
        state = agent.run(self.root, str(self.intake), converse=breaks)
        self.assertIn("gateway down", state.stopped_because)
        self.assertEqual(state.gaps_now, state.gaps_at_start)

    def test_a_failing_tool_is_information_rather_than_a_stop(self):
        state = agent.run(self.root, str(self.intake),
                          converse=_scripted(_call("read_document", name="nothing.pdf")))
        self.assertGreaterEqual(state.calls, 1)
        self.assertIn("read_document", state.tools_run)


class TestTheToolSurface(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.intake = self.root / "intake.xlsx"
        shutil.copy(EXAMPLES / "4_spans_drawn_wrongly.xlsx", self.intake)
        self.tools = self._tools()

    def _tools(self):
        pack = agent.Pack(self.root)
        return {t.name: t for t in agent.build_tools(pack, str(self.intake), {}, {})}

    def test_it_lists_what_was_submitted(self):
        sources = (self.root / "sources" / "model_doc")
        sources.mkdir(parents=True)
        (sources / "notes.md").write_text("# Agent\n\nIt verifies then files.\n")
        self.assertIn("notes.md", self._tools()["list_sources"].run())

    def test_a_document_comes_back_with_its_page_markers(self):
        sources = (self.root / "sources" / "model_doc")
        sources.mkdir(parents=True)
        (sources / "notes.md").write_text("# Agent\n\nIt verifies the cardmember.\n")
        read = self._tools()["read_document"].run(name="notes.md")
        self.assertIn("verifies the cardmember", read)
        self.assertIn("[", read, "no locator, so nothing can be traced back to a page")

    def test_it_says_so_plainly_when_nothing_was_submitted(self):
        """Any combination of documents and diagrams is a pack, including none of either."""
        self.assertIn("Nothing was submitted", self.tools["list_sources"].run())

    def test_asking_for_a_file_that_is_not_there_is_answered_not_raised(self):
        self.assertIn("no submitted document",
                      self.tools["read_document"].run(name="absent.pdf"))

    def test_the_audit_is_the_declaration_rather_than_the_documents(self):
        reported = self.tools["audit_declaration"].run()
        self.assertIn("CAP-02", reported)

    def test_what_is_declared_renders_the_graph_with_its_edges(self):
        rendered = self.tools["what_is_declared"].run()
        self.assertIn("DEC-01", rendered)
        self.assertIn("Capabilities", rendered)


class TestQuestionsHaveToBeAnswerable(unittest.TestCase):
    """The failure is not asking too many questions. It is asking unanswerable ones -- a list of
    forty "is this clear?" is a list nobody opens."""

    def setUp(self):
        self.problems = ["DEC-07: outcome \"Timeout\" leads to no declared state",
                         "CAP-02: which states does it hand on at?"]
        self.recorded = []
        self.tool = agent.build_question_tool(self.problems, self.recorded)

    def test_a_question_naming_a_row_the_audit_raised_is_recorded(self):
        self.tool.run(question="DEC-07 times out and leads nowhere — where does it go?")
        self.assertEqual(len(self.recorded), 1)

    def test_a_vague_question_is_refused_with_the_reason(self):
        answer = self.tool.run(question="Is decision seven clear?")
        self.assertEqual(self.recorded, [])
        self.assertIn("does not name a row", answer)

    def test_a_question_about_something_nobody_flagged_is_refused(self):
        self.tool.run(question="What is the retry limit on DEC-99, which nobody mentioned?")
        self.assertEqual(self.recorded, [])

    def test_the_same_question_twice_is_recorded_once(self):
        for _ in range(2):
            self.tool.run(question="CAP-02 hands on at which states?")
        self.assertEqual(len(self.recorded), 1)

    def test_questions_reach_the_result(self):
        root = Path(tempfile.mkdtemp())
        intake = root / "intake.xlsx"
        shutil.copy(EXAMPLES / "4_spans_drawn_wrongly.xlsx", intake)
        state = agent.run(root, str(intake), converse=_scripted(
            _call("ask_the_model_owner", question="CAP-02 hands on at which states?")))
        self.assertEqual(state.questions, ["CAP-02 hands on at which states?"])


class TestTheOracle(unittest.TestCase):
    def test_it_names_an_outcome_that_leads_nowhere(self):
        from openpyxl import load_workbook
        root = Path(tempfile.mkdtemp())
        intake = root / "intake.xlsx"
        shutil.copy(EXAMPLES / "1_disputes_three_blocks.xlsx", intake)
        book = load_workbook(intake)
        # Give DEC-01 an outcome nothing declares a destination for.
        for row in book["L3 Decisions"].iter_rows(min_row=2):
            if row[0].value == "DEC-01":
                row[4].value = "Dispute / Other intent / Unclear"
        book.save(intake)

        problems = agent.outstanding(str(intake))
        self.assertTrue(any("Unclear" in p and "DEC-01" in p for p in problems), problems)

    def test_an_unreadable_declaration_is_reported_rather_than_raised(self):
        broken = Path(tempfile.mkdtemp()) / "not-a-workbook.xlsx"
        broken.write_text("this is not a workbook")
        self.assertTrue(agent.outstanding(str(broken))[0].startswith("The declaration could not"))


if __name__ == "__main__":
    unittest.main()


class TestTheSwitch(unittest.TestCase):
    """Off by default, and never able to fail the stage when on.

    The fixed sequence works, is what every other test exercises, and needs nothing from the
    gateway beyond an ordinary completion. The loop needs tool-calling, which not every gateway
    offers -- so it is a switch, and a gateway that cannot run it leaves a drafted declaration
    rather than a failed stage.
    """

    def test_it_is_off_unless_switched_on(self):
        import os
        from metric.llm import config
        os.environ.pop("LLM_INTAKE_LOOP", None)
        self.assertFalse(config.intake_loop())

    def test_a_gateway_that_cannot_run_it_falls_back_to_the_sequence(self):
        """An intake stage producing nothing because tool-calling is unavailable would be worse
        than one costing a few more calls."""
        from unittest import mock

        from metric.web import runners

        root = Path(tempfile.mkdtemp())
        (root / "sources" / "model_doc").mkdir(parents=True)
        (root / "sources" / "model_doc" / "notes.md").write_text("# Agent\n\nIt verifies.\n")
        target = root / "drafted_intake.xlsx"

        class _Workspace:
            def __init__(self, root):
                self.root = root
            def note_lines(self):
                return []
            def is_marked_for_redaction(self, group, name):
                return False
            def artifact_path(self, *_):
                return None
            def state(self, _):
                return type("S", (), {"artifacts": {}})()

        fell_back = {}

        def _sequence(workspace, provided, target, context, ours, report, cancel, summary):
            fell_back["yes"] = True
            summary["_action"] = "drafted"

        with mock.patch("metric.phases.intake.intake.agent.run",
                        side_effect=RuntimeError("no tool calling here")), \
             mock.patch.object(runners, "_needs_reading", return_value=False), \
             mock.patch.object(runners, "_draft_or_revise", _sequence):
            result = runners._read_and_draft_with_loop(_Workspace(root), target,
                                                       lambda *a, **k: None)

        self.assertTrue(fell_back, "the stage produced nothing rather than falling back")
        self.assertIn("could not run", str(result["The loop"]))

    def test_the_summary_says_what_it_did(self):
        root = Path(tempfile.mkdtemp())
        intake = root / "intake.xlsx"
        shutil.copy(EXAMPLES / "4_spans_drawn_wrongly.xlsx", intake)
        state = agent.run(root, str(intake), converse=_scripted(_call("audit_declaration")))
        summary = state.summary()
        self.assertIn("Gaps at the start", summary)
        self.assertIn("Stopped because", summary)
        self.assertGreaterEqual(summary["Model calls"], 1)


_DECLARATION = {
    "use_case": {"name": "Disputes assistant", "objective": "Resolve disputed charges"},
    "personas": [{"id": "P1", "name": "Cardmember", "applies_to": "Wants a charge investigated",
                  "is_default": True}],
    "capabilities": [{"id": "CAP-01", "name": "Identification", "type": "Gating"}],
    "decisions": [{"id": "DEC-01", "name": "Identify the cardmember", "capability_id": "CAP-01",
                   "inputs": "card details", "outcomes": ["Identified", "Cannot identify"],
                   "input_source": "User", "max_attempts": 2, "outcome_condition": ""}],
    "states": [
        {"id": "S-00", "reached_via": "Start", "description": "The chat opens",
         "next_decisions": ["DEC-01"], "is_terminal": False, "outcome_type": ""},
        {"id": "S-01", "reached_via": "DEC-01=Identified", "description": "Identified",
         "next_decisions": [], "is_terminal": True, "outcome_type": "Happy path"},
        {"id": "S-02", "reached_via": "DEC-01=Cannot identify", "description": "Locked out",
         "next_decisions": [], "is_terminal": True, "outcome_type": "Termination"}],
    "tools": [{"name": "Identity service", "capability_id": "CAP-01", "changes_state": False}],
    "confidence": {}, "review_notes": [],
}


class TestItBuildsADeclarationFromNothing(unittest.TestCase):
    """The loop replaces the sequence, so it has to produce a declaration rather than repair one."""

    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        (self.root / "sources" / "model_doc").mkdir(parents=True)
        (self.root / "sources" / "model_doc" / "notes.md").write_text(
            "# Disputes assistant\n\nIt identifies the cardmember, then files a dispute.\n")
        self.intake = self.root / "drafted_intake.xlsx"

    def test_a_written_declaration_lands_in_a_readable_workbook(self):
        import json as _json
        from metric.domain.intake import read_intake

        state = agent.run(self.root, str(self.intake), converse=_scripted(
            _call("write_declaration", declaration=_json.dumps(_DECLARATION))))
        self.assertTrue(self.intake.exists())
        intake = read_intake(str(self.intake))
        # Two gaps remain, and they are the right two: this declaration states no agent type
        # and no success criteria, and the audit is what says so.
        self.assertEqual([d.id for d in intake.decisions], ["DEC-01"])
        self.assertTrue(all("use_case" in p for p in agent.outstanding(str(self.intake))))

    def test_it_costs_far_less_than_the_sequence_it_replaces(self):
        """Seven to nine calls were spent before this loop existed, every run, whatever the pack.
        A one-document pack that resolves in two turns must not cost more than that."""
        import json as _json
        state = agent.run(self.root, str(self.intake), converse=_scripted(
            _call("read_document", name="notes.md"),
            _call("write_declaration", declaration=_json.dumps(_DECLARATION))))
        self.assertLessEqual(state.calls, 4, "more than the sequence it replaces")

    def test_what_the_documents_establish_is_written_for_later_stages(self):
        """Every stage after this one sees the graph and not the documents, so the reading has to
        land on disk in the two files they read."""
        import json as _json
        from metric.phases.intake.intake.evidence import EvidenceRecord, FacetAnswer

        record = EvidenceRecord(answers=[FacetAnswer(
            facet="use_case", answer="It resolves disputed card charges without a person.",
            points=["It resolves disputed card charges without a person."],
            confidence="High")])

        def reads(paths, **kwargs):
            return record

        with mock.patch("metric.phases.intake.intake.extraction.extract_documents", reads):
            agent.run(self.root, str(self.intake), converse=_scripted(
                _call("read_the_pack"),
                _call("write_declaration", declaration=_json.dumps(_DECLARATION))))

        context = (self.root / "ingest_context.md").read_text()
        self.assertIn("resolves disputed card charges", context)
        self.assertTrue((self.root / "ingest_evidence.json").exists())

    def test_the_reading_lands_even_when_the_loop_never_asks_for_it(self):
        """The defect this guards. A loop that opened files one at a time with read_document
        produced no record, so no context document was written -- and every stage after this one
        is grounded on that file rather than on the documents, which they never see."""
        import json as _json
        from metric.phases.intake.intake.evidence import EvidenceRecord, FacetAnswer

        record = EvidenceRecord(answers=[FacetAnswer(
            facet="use_case", answer="It resolves disputed card charges without a person.",
            points=["It resolves disputed card charges without a person."],
            confidence="High")])

        with mock.patch("metric.phases.intake.intake.extraction.extract_documents",
                        lambda paths, **kwargs: record):
            agent.run(self.root, str(self.intake), converse=_scripted(
                _call("read_document", name="notes.md"),
                _call("write_declaration", declaration=_json.dumps(_DECLARATION))))

        self.assertIn("resolves disputed card charges",
                      (self.root / "ingest_context.md").read_text())
        self.assertTrue((self.root / "ingest_evidence.json").exists())

    def test_the_reading_is_in_front_of_the_loop_before_its_first_turn(self):
        """A turn spent asking for something needed on every run is a turn wasted."""
        from metric.phases.intake.intake.evidence import EvidenceRecord, FacetAnswer

        record = EvidenceRecord(answers=[FacetAnswer(
            facet="use_case", answer="It resolves disputed card charges without a person.",
            points=["It resolves disputed card charges without a person."],
            confidence="High")])
        seen = {}

        def converse(messages):
            seen.setdefault("first", list(messages))
            return {"content": "", "tool_calls": []}

        with mock.patch("metric.phases.intake.intake.extraction.extract_documents",
                        lambda paths, **kwargs: record):
            agent.run(self.root, str(self.intake), converse=converse)

        opening = seen["first"][-1]["content"]
        self.assertIn("resolves disputed card charges", opening)
        self.assertIn("notes.md", opening, "the file list went missing with the reading added")

    def test_the_pack_is_read_with_the_pipeline_built_for_it(self):
        """Not re-implemented in the loop. Reading documents one at a time as raw text threw away
        the grounding, the facet coverage, and -- worst -- the multi-image diagram passes."""
        seen = {}

        def reads(paths, **kwargs):
            from metric.phases.intake.intake.evidence import EvidenceRecord
            seen["paths"] = list(paths)
            seen["kwargs"] = kwargs
            return EvidenceRecord()

        with mock.patch("metric.phases.intake.intake.extraction.extract_documents", reads):
            agent.run(self.root, str(self.intake), converse=_scripted(_call("read_the_pack")),
                      diagram_mode="same_flow")

        self.assertEqual([p.name for p in seen["paths"]], ["notes.md"])
        self.assertEqual(seen["kwargs"]["diagram_mode"], "same_flow",
                         "the mode somebody chose for their images was not passed through")
        self.assertIn("should_redact", seen["kwargs"])

    def test_the_pack_is_read_once_however_often_it_is_asked_for(self):
        """The most expensive call available. Asking twice returns the reading, not a second one."""
        calls = []

        def reads(paths, **kwargs):
            from metric.phases.intake.intake.evidence import EvidenceRecord
            calls.append(paths)
            return EvidenceRecord()

        with mock.patch("metric.phases.intake.intake.extraction.extract_documents", reads):
            read_once = {}
            tools = {t.name: t for t in agent.build_tools(
                agent.Pack(self.root), str(self.intake), read_once, {})}
            first = tools["read_the_pack"].run()
            self.assertEqual(tools["read_the_pack"].run(), first)
            self.assertEqual(len(calls), 1, "the pack was read a second time")


class TestWhatWriteDeclarationRefuses(unittest.TestCase):
    """The one tool that changes anything, so the one that has to refuse. An empty workbook
    overwriting a partial one is the worst thing a turn could do, and the easiest to produce."""

    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.intake = self.root / "intake.xlsx"
        self.tools = {t.name: t for t in agent.build_tools(
            agent.Pack(self.root), str(self.intake), {}, {})}

    def test_something_that_is_not_json_is_refused_with_the_reason(self):
        answer = self.tools["write_declaration"].run(declaration="here is the graph, roughly")
        self.assertIn("not written", answer.lower())
        self.assertFalse(self.intake.exists())

    def test_a_declaration_with_no_graph_in_it_is_refused(self):
        import json as _json
        answer = self.tools["write_declaration"].run(
            declaration=_json.dumps({"use_case": {"name": "x"}, "decisions": [], "states": []}))
        self.assertIn("describes no agent", answer)
        self.assertFalse(self.intake.exists())

    def test_a_write_reports_what_is_still_outstanding(self):
        import json as _json
        thin = dict(_DECLARATION,
                    decisions=[dict(_DECLARATION["decisions"][0],
                                    outcomes=["Identified", "Cannot identify", "Timed out"])])
        answer = self.tools["write_declaration"].run(declaration=_json.dumps(thin))
        self.assertIn("Still outstanding", answer)
        self.assertIn("Timed out", answer)


class TestSpansSurviveTheLoop(unittest.TestCase):
    """Where a block begins and ends is a person's judgement about the agent and decides how the
    whole scenario space is enumerated. The loop fills in everything else about a capability."""

    def test_a_span_already_drawn_is_carried_across_a_write(self):
        import json as _json
        from metric.domain.intake import read_intake

        root = Path(tempfile.mkdtemp())
        intake = root / "intake.xlsx"
        shutil.copy(EXAMPLES / "1_disputes_three_blocks.xlsx", intake)
        before = {c.id: (c.entry_states, c.exit_states)
                  for c in read_intake(str(intake)).capabilities}

        tools = {t.name: t for t in agent.build_tools(agent.Pack(root), str(intake), {}, {})}
        # The loop rewrites the declaration and says nothing about spans, as it is told not to.
        stripped = dict(_DECLARATION, capabilities=[
            {"id": cid, "name": f"Renamed {cid}", "type": "Gating"} for cid in before])
        tools["write_declaration"].run(declaration=_json.dumps(stripped))

        after = read_intake(str(intake))
        self.assertEqual({c.id: (c.entry_states, c.exit_states) for c in after.capabilities},
                         before)
        self.assertTrue(any(c.name.startswith("Renamed") for c in after.capabilities),
                        "the loop's own corrections were lost along with the spans")


class TestTheCapabilityLifecycle(unittest.TestCase):
    """A capability is filled in by two hands, and each owns a different part of it.

    A person draws the span -- where the block starts and where it hands on -- against the
    drawing, because that decides how the whole scenario space is enumerated and no model can read
    it off a document. Everything else about the capability is the loop's: its name, its type,
    what it does, from the decisions inside the span and the documents describing them.

    The failure worth guarding is the loop overwriting the half it does not own.
    """

    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.intake = self.root / "intake.xlsx"
        shutil.copy(EXAMPLES / "1_disputes_three_blocks.xlsx", self.intake)

    def _declaration(self, capabilities):
        from metric.domain.intake import read_intake
        current = read_intake(str(self.intake))
        return {
            "use_case": {"name": current.name, "objective": "Resolve disputed charges"},
            "personas": [{"id": p.id, "name": p.name, "applies_to": " ".join(p.applies_to),
                          "is_default": p.is_default} for p in current.personas],
            "capabilities": capabilities,
            "decisions": [{"id": d.id, "name": d.name, "capability_id": d.trigger_capability,
                           "inputs": d.inputs, "outcomes": list(d.variants),
                           "input_source": d.input_source, "max_attempts": d.max_attempts,
                           "outcome_condition": d.outcome_condition} for d in current.decisions],
            "states": [{"id": s.id, "reached_via": s.reached_via, "description": s.description,
                        "next_decisions": list(s.next_decisions), "is_terminal": s.is_terminal,
                        "outcome_type": s.outcome_type} for s in current.states],
            "tools": [{"name": t.name, "capability_id": t.capability_id,
                       "changes_state": t.state_changing} for t in current.tools],
            "confidence": {}, "review_notes": []}

    def _write(self, capabilities):
        import json as _json
        tools = {t.name: t for t in agent.build_tools(
            agent.Pack(self.root), str(self.intake), {}, {})}
        return tools["write_declaration"].run(
            declaration=_json.dumps(self._declaration(capabilities)))

    def test_a_span_the_user_redrew_survives_the_loop_rewriting_everything_else(self):
        from metric.domain.intake import read_intake, set_capability_span

        set_capability_span(str(self.intake), "CAP-02", ["S-02"], ["S-07", "S-08"])
        self._write([{"id": "CAP-01", "name": "Cardmember identification", "type": "Gating"},
                     {"id": "CAP-02", "name": "Cardmember verification", "type": "Gating"},
                     {"id": "CAP-03", "name": "Charge handling", "type": "Transactional"}])

        after = {c.id: c for c in read_intake(str(self.intake)).capabilities}
        self.assertEqual(after["CAP-02"].entry_states, ("S-02",), "the redrawn span was lost")
        self.assertEqual(after["CAP-02"].exit_states, ("S-07", "S-08"))
        self.assertEqual(after["CAP-02"].name, "Cardmember verification",
                         "the loop's enrichment was lost along with it")

    def test_the_loop_fills_in_a_type_the_workbook_left_blank(self):
        """An untyped capability drops its adversarial probes silently, which is exactly the kind
        of blank the loop exists to close."""
        from metric.domain.intake import read_intake

        self._write([{"id": "CAP-01", "name": "Identification", "type": ""},
                     {"id": "CAP-02", "name": "Verification", "type": ""},
                     {"id": "CAP-03", "name": "Charge handling", "type": "Transactional"}])
        self.assertEqual(
            [c.type for c in read_intake(str(self.intake)).capabilities if c.id == "CAP-01"], [""])

        self._write([{"id": "CAP-01", "name": "Identification", "type": "Gating"},
                     {"id": "CAP-02", "name": "Verification", "type": "Gating"},
                     {"id": "CAP-03", "name": "Charge handling", "type": "Transactional"}])
        typed = {c.id: c.type for c in read_intake(str(self.intake)).capabilities}
        self.assertEqual(typed["CAP-01"], "Gating")

    def test_redrawing_a_span_changes_what_is_enumerated(self):
        """The whole reason the span is the person's to draw."""
        from metric.domain.intake import read_intake, set_capability_span
        from metric.pipeline import build_scenario_space

        before = len(build_scenario_space(read_intake(str(self.intake)), with_probes=False))
        set_capability_span(str(self.intake), "CAP-02", ["S-02"], ["S-07", "S-08"])
        after = len(build_scenario_space(read_intake(str(self.intake)), with_probes=False))
        self.assertEqual((before, after), (14, 11))

    def test_the_loop_folds_a_capability_declared_twice(self):
        """Two capabilities that are one capability are two blocks of the space where there is
        one, so this matters more than tidiness."""
        from metric.domain.intake import read_intake

        self._write([{"id": "CAP-01", "name": "Identification", "type": "Gating"},
                     {"id": "CAP-02", "name": "Identification service", "type": "Gating"},
                     {"id": "CAP-03", "name": "Charge handling", "type": "Transactional"}])
        kept = [c.id for c in read_intake(str(self.intake)).capabilities]
        self.assertEqual(kept, ["CAP-01", "CAP-03"])


class TestDiagramsAreReadTheWayTheyWereBuiltToBe(unittest.TestCase):
    """Several images are three passes, not one call per picture.

    Reading them one at a time always took the single-image path, whose prompt says the image *is*
    the whole flow. Three images each got told that, each numbered its boxes from one, and the
    collisions were dropped when the readings were merged — whole pictures vanished silently, and
    what survived was a graph welded out of three drawings that each thought they were complete.
    """

    def setUp(self):
        import base64

        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")
        self.root = Path(tempfile.mkdtemp())
        (self.root / "sources" / "diagrams").mkdir(parents=True)
        for n in (1, 2, 3):
            (self.root / "sources" / "diagrams" / f"flow{n}.png").write_bytes(png)
        self.intake = self.root / "i.xlsx"
        self.prompts = []

    def _read(self, **kwargs):
        import json as _json

        def images(system, user, images=None, **k):
            self.prompts.append(("image", user))
            return _json.dumps({"nodes": [{"ref": "n1", "kind": "decision", "label": "Check"}],
                                "edges": [], "continues_offpage": [], "observations": []})

        def text(system, user, **k):
            self.prompts.append(("text", user))
            return _json.dumps({"capabilities": [], "decisions": [], "states": [],
                                "observations": [], "unresolved": [], "answer": "",
                                "points": [], "unknowns": [], "sources": [],
                                "confidence": "Low"})

        with mock.patch("metric.phases.intake.intake.extraction.ask_llm", text), \
             mock.patch("metric.phases.intake.intake.extraction.ask_llm_with_images", images):
            return agent.run(self.root, str(self.intake),
                             converse=_scripted(_call("read_the_pack")), **kwargs)

    def test_no_image_is_told_it_is_the_whole_flow(self):
        """The specific regression. That prompt is correct for one image and a lie for three."""
        self._read()
        told = [p for kind, p in self.prompts if "It is the only image in the pack" in p]
        self.assertEqual(told, [])

    def test_each_image_is_read_on_its_own_and_the_result_checked_against_them(self):
        """Three reads and a look-again. The last one goes back to the pictures with whatever the
        joined graph could not account for, which is where a missed arrow gets caught."""
        self._read()
        self.assertEqual(sum(1 for kind, _ in self.prompts if kind == "image"), 4)

    def test_and_then_they_are_joined(self):
        """The pass that decides whether a box in one picture is the same box in another."""
        self._read()
        joined = [p for kind, p in self.prompts if kind == "text" and "IMAGE 1 of 3" in p]
        self.assertEqual(len(joined), 1)

    def test_the_mode_somebody_chose_for_their_images_is_honoured(self):
        """Stitching two drawings of one flow welds them end to end and invents routes the agent
        does not have. Reconciling pieces of one picture loses a step. The person knows which."""
        self._read(diagram_mode="same_flow")
        joined = next(p for kind, p in self.prompts if kind == "text" and "IMAGE 1 of 3" in p)
        self.assertIn("same flow", joined.lower())
