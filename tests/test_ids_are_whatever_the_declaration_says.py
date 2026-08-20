"""An id is whatever the declaration calls it, not what this tool would have called it.

Every place that has to pick an id out of a cell or a sentence used to look for a prefix and a
*number* -- ``S-\\d+``, ``DEC-\\d+``. That is the convention this tool writes, and it is a
reasonable one, but it is not a convention it may impose on a declaration it is only reading. A
workflow whose opening state is called ``S-START`` was legible to every human who looked at it and
invisible to the parser.

The failure was silent and looked like something else entirely. Ticking ``S-START`` as a
capability's entry state wrote it to the workbook correctly; reading the workbook back found no
state ids in that cell at all; the page redrew from what was read; and the tick was gone. The
report was "it will not save this state", which is exactly what it looks like from outside and
nowhere near where the fault was.
"""
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from metric.domain import editing
from metric.domain.graph import DecisionGraph
from metric.domain.intake import read_intake
from metric.shared.text import parse_reached_via

EXAMPLE = (Path(__file__).resolve().parent.parent
           / "examples" / "intakes" / "2_travel_no_spans.xlsx")


def _renamed(path: Path, old: str, new: str) -> None:
    """The same declaration with one state id written the way a person would write it."""
    book = load_workbook(path)
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and old in cell.value:
                    cell.value = cell.value.replace(old, new)
    book.save(path)


class TestAStateNamedForWhatItIs(unittest.TestCase):
    def setUp(self):
        self.path = Path(tempfile.mkdtemp()) / "intake.xlsx"
        shutil.copy(EXAMPLE, self.path)
        _renamed(self.path, "S-00", "S-START")

    def test_the_state_is_read(self):
        self.assertIn("S-START", [s.id for s in read_intake(str(self.path)).states])

    def test_it_is_still_where_the_graph_starts(self):
        intake = read_intake(str(self.path))
        graph = DecisionGraph(intake.decisions, intake.states)
        self.assertEqual(graph.start_states, ["S-START"])

    def test_a_span_entered_there_survives_the_round_trip(self):
        """The reported symptom, end to end: save it, read it back, and it is still there."""
        editing.apply_edits(str(self.path), [
            {"kind": "capability", "key": "CAP-01", "action": "upsert",
             "fields": {"entry_states": ["S-START"], "exit_states": ["S-03"]}}])
        found = next(c for c in read_intake(str(self.path)).capabilities if c.id == "CAP-01")
        self.assertEqual(found.entry_states, ("S-START",))
        self.assertTrue(found.is_bounded)

    def test_the_capability_is_walked_as_a_block(self):
        """A span that reads back empty is a capability that contributes no scenarios, which is
        the consequence that matters and the one nobody would connect to an id shape."""
        from metric.domain.graph import spans_for
        editing.apply_edits(str(self.path), [
            {"kind": "capability", "key": "CAP-01", "action": "upsert",
             "fields": {"entry_states": ["S-START"], "exit_states": ["S-03"]}}])
        intake = read_intake(str(self.path))
        graph = DecisionGraph(intake.decisions, intake.states)
        spans = spans_for(graph, intake.capabilities)
        self.assertIn("CAP-01", [s.capability_id for s in spans])


class TestADecisionNamedForWhatItDoes(unittest.TestCase):
    def setUp(self):
        self.path = Path(tempfile.mkdtemp()) / "intake.xlsx"
        shutil.copy(EXAMPLE, self.path)
        _renamed(self.path, "DEC-01", "DEC-FIND")

    def test_a_state_still_offers_it(self):
        intake = read_intake(str(self.path))
        offering = next(s for s in intake.states if "DEC-FIND" in s.next_decisions)
        self.assertTrue(offering.id)

    def test_the_outcome_that_reaches_a_state_is_still_read(self):
        self.assertEqual(parse_reached_via("DEC-FIND=Found"), [("DEC-FIND", "Found")])

    def test_the_graph_connects(self):
        intake = read_intake(str(self.path))
        graph = DecisionGraph(intake.decisions, intake.states)
        self.assertTrue(graph.states_offering("DEC-FIND"))
        self.assertTrue(graph.successor("DEC-FIND", "Found"))


class TestScrapingPrefersWhatWasDeclared(unittest.TestCase):
    """The shape is a fallback, not the rule -- which is what makes any naming scheme work."""

    def setUp(self):
        from metric.domain.intake import _STATE_TOKEN
        self.shape = _STATE_TOKEN

    def _named(self, cell, declared):
        from metric.domain.intake import _named
        return _named(cell, declared, self.shape)

    def test_any_separator_reads_the_same(self):
        for cell in ("S-START, S-04", "S-START and S-04", "S-START\nS-04", "S-START; S-04"):
            self.assertEqual(self._named(cell, ["S-START", "S-04"]), ("S-START", "S-04"), cell)

    def test_ids_with_no_prefix_at_all_are_read(self):
        """Nothing about a span requires this tool's naming. A declaration that calls its states
        WELCOME and GOODBYE is naming them better than S-00 and S-11 do."""
        self.assertEqual(self._named("WELCOME, GOODBYE", ["WELCOME", "GOODBYE"]),
                         ("WELCOME", "GOODBYE"))

    def test_a_longer_id_is_not_eaten_by_a_shorter_one(self):
        self.assertEqual(self._named("S-10", ["S-1", "S-10"]), ("S-10",))

    def test_the_declared_spelling_wins(self):
        """One id written two ways is one id, and which spelling the walk keys on cannot depend on
        which cell happened to be read."""
        self.assertEqual(self._named("s-start", ["S-START"]), ("S-START",))

    def test_a_reference_to_a_state_that_does_not_exist_is_still_read(self):
        """So the audit can report it as dangling. Dropping it here is how a typo becomes a span
        that silently covers half the block it names."""
        self.assertEqual(self._named("S-99", ["S-01"]), ("S-99",))

    def test_nothing_in_an_empty_cell(self):
        self.assertEqual(self._named("   ", ["S-01"]), ())


if __name__ == "__main__":
    unittest.main()
