"""Asking which column is which, instead of guessing from the heading.

Heading matching fails in the way that hurts most: silently and plausibly. A column headed
`Utterance` is found; one headed `cust_msg_body_txt` is not, and the reader falls back to "the
widest column carries the words" -- which on a real export picks the agent's reasoning trace and
files it under the user. Nothing about the result looks wrong. The coverage figure is simply built
on the wrong column, and the person defending it has no way to see that.

So the layout is asked about once, over the header and a handful of rows, and every row is then
read in code from what came back. One call for a file of any size, because the model reads the
header rather than the transcripts.

Everything here is also about what happens when that call is no help. A reply that will not parse,
a column named that the file does not have, a layout word this does not know, no gateway at all --
every one of them has to end in a reading rather than an error, because a heading-matched figure
whose provenance is on the page beats no figure at all.
"""
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from scenario_generator.ingest import conversation_migration
from scenario_generator.ingest.conversations import read_conversations

_HEADER = ["sess_key", "seq_no", "party_cd", "cust_msg_body_txt", "agnt_rsn_trace", "tc_ref"]
_USER_LINE = "I do not recognise a charge on my statement at all."
_AGENT_LINE = "Can you confirm the last four digits of your SSN please?"

_LAYOUT = {"layout": "turn_per_row", "conversation_id": "sess_key", "speaker": "party_cd",
           "text": "cust_msg_body_txt", "user_text": "", "agent_text": "",
           "user_values": ["C"], "agent_values": ["A"], "scenario_id": "", "group": "tc_ref",
           "note": "party_cd is C for the customer and A for the agent."}


def _export(directory: Path, conversations: int = 3) -> Path:
    """An export shaped the way a real logging system produces one."""
    book = Workbook()
    sheet = book.active
    sheet.append(_HEADER)
    for n in range(1, conversations + 1):
        # The trace is deliberately the widest column: that is what "widest wins" picks.
        sheet.append([f"K{n}", 1, "C", _USER_LINE, "internal reasoning " * 12, "TC-77"])
        sheet.append([f"K{n}", 2, "A", _AGENT_LINE, "internal reasoning " * 12, "TC-77"])
    path = directory / "export.xlsx"
    book.save(path)
    return path


def _replying(payload) -> tuple:
    """A completion returning ``payload``, and the list of prompts it was sent."""
    seen = []

    def complete(system, user, **kwargs):
        seen.append(user)
        return payload if isinstance(payload, str) else json.dumps(payload)

    return complete, seen


class TestItReadsWhatTheHeadingsCannot(unittest.TestCase):
    def setUp(self):
        self.path = _export(Path(tempfile.mkdtemp()))

    def test_by_headings_alone_the_rows_never_group_into_conversations(self):
        """The state this replaces, pinned so the improvement is not imaginary."""
        conversations, _ = read_conversations(self.path, migrate=False)
        self.assertEqual(len(conversations), 6)

    def test_asked_about_the_layout_it_groups_correctly(self):
        complete, _ = _replying(_LAYOUT)
        conversations, _ = read_conversations(self.path, complete=complete)
        self.assertEqual(len(conversations), 3)
        self.assertEqual(len(conversations[0].turns), 2)

    def test_it_takes_the_column_it_was_told_rather_than_the_widest(self):
        complete, _ = _replying(_LAYOUT)
        conversations, _ = read_conversations(self.path, complete=complete)
        self.assertIn(_USER_LINE, conversations[0].transcript)
        self.assertNotIn("internal reasoning", conversations[0].transcript)

    def test_a_speaker_code_the_word_list_cannot_know_is_translated(self):
        """`C` and `A` are the normal case in an export, not the exception."""
        complete, _ = _replying(_LAYOUT)
        conversations, _ = read_conversations(self.path, complete=complete)
        self.assertEqual([t.speaker for t in conversations[0].turns], ["user", "agent"])
        self.assertEqual(len(conversations[0].user_turns), 1)

    def test_the_teams_own_label_is_carried_but_kept_as_a_label(self):
        complete, _ = _replying(_LAYOUT)
        conversations, _ = read_conversations(self.path, complete=complete)
        self.assertEqual(conversations[0].group, "TC-77")
        self.assertEqual(conversations[0].scenario_id, "")

    def test_how_it_was_read_says_a_call_worked_the_columns_out(self):
        """A misread column halves a coverage figure quietly, so the provenance is on the page."""
        complete, _ = _replying(_LAYOUT)
        _, how = read_conversations(self.path, complete=complete)
        self.assertIn("columns worked out by a first pass", how)
        self.assertIn("C for the customer", how)


class TestItCostsOneCallWhateverTheFileSize(unittest.TestCase):
    def test_one_call_for_a_file_of_any_size(self):
        directory = Path(tempfile.mkdtemp())
        complete, seen = _replying(_LAYOUT)
        read_conversations(_export(directory, conversations=200), complete=complete)
        self.assertEqual(len(seen), 1)

    def test_the_call_is_shown_the_header_and_a_sample_rather_than_the_file(self):
        directory = Path(tempfile.mkdtemp())
        complete, seen = _replying(_LAYOUT)
        read_conversations(_export(directory, conversations=200), complete=complete)
        self.assertIn("sess_key", seen[0])
        self.assertLessEqual(seen[0].count("K1"), 4)
        self.assertNotIn("K150", seen[0])


class TestWhenTheCallIsNoHelp(unittest.TestCase):
    """Every one of these has to end in a reading, never in an error."""

    def setUp(self):
        self.path = _export(Path(tempfile.mkdtemp()))

    def _falls_back(self, payload):
        complete, _ = _replying(payload)
        conversations, how = read_conversations(self.path, complete=complete)
        self.assertTrue(conversations)
        self.assertNotIn("columns worked out", how)
        return conversations

    def test_a_reply_that_will_not_parse(self):
        self._falls_back("not json at all")

    def test_a_call_that_raises(self):
        def complete(system, user, **kwargs):
            raise RuntimeError("gateway down")

        conversations, how = read_conversations(self.path, complete=complete)
        self.assertTrue(conversations)
        self.assertNotIn("columns worked out", how)

    def test_a_column_the_file_does_not_have(self):
        """More likely a misread sample than a discovery, and reading a column that is not there
        is worse than reading the wrong one by heading."""
        self._falls_back(dict(_LAYOUT, text="message_body"))

    def test_a_layout_word_this_does_not_know(self):
        self._falls_back(dict(_LAYOUT, layout="threaded"))

    def test_a_reply_naming_no_column_that_carries_the_words(self):
        self._falls_back(dict(_LAYOUT, text="", user_text="", agent_text=""))

    def test_no_gateway_at_all(self):
        conversations, how = read_conversations(self.path)
        self.assertTrue(conversations)
        self.assertNotIn("columns worked out", how)


class TestTheOtherTwoReadingsAreUntouched(unittest.TestCase):
    def test_our_own_template_never_reaches_the_call(self):
        """It states which scenario each conversation was run against, so there is nothing to
        work out and nothing to pay for."""
        from scenario_generator.core.models import (Capability, Decision, IntakeData, Persona,
                                                    State, Tool)
        from scenario_generator.io import write_data_template
        from scenario_generator.pipeline import build_scenario_space
        from openpyxl import load_workbook

        intake = IntakeData(
            use_case={"Use case name": "Disputes", "Business objective": "Resolve"},
            personas=[Persona("P1", "Cardmember", [], True)],
            capabilities=[Capability("CAP-01", "Identity", "Gating")],
            decisions=[Decision("DEC-01", "Identity", "CAP-01", "", ["Pass", "Fail"])],
            states=[State("S-00", "Start", "Start", ["DEC-01"], False),
                    State("S-01", "DEC-01=Pass", "Verified", [], True, "Happy path"),
                    State("S-02", "DEC-01=Fail", "Locked out", [], True, "Termination")],
            tools=[Tool("Identity service", "CAP-01", True)])
        scenarios = build_scenario_space(intake, with_probes=False)
        for scenario in scenarios:
            scenario.turn_plan = "1. Open."
        directory = Path(tempfile.mkdtemp())
        template = directory / "data_template.xlsx"
        write_data_template(str(template), intake, scenarios)

        book = load_workbook(template)
        log = book["Variation_Log"]
        for row in range(2, log.max_row + 1):
            log.cell(row=row, column=4).value = _USER_LINE
            log.cell(row=row, column=5).value = _AGENT_LINE
        book.save(template)

        complete, seen = _replying(_LAYOUT)
        conversations, how = read_conversations(template, complete=complete)
        self.assertEqual(seen, [])
        self.assertIn("data template returned filled in", how)
        self.assertTrue(all(c.scenario_id for c in conversations))

    def test_prose_with_no_table_never_reaches_the_call(self):
        directory = Path(tempfile.mkdtemp())
        path = directory / "transcripts.md"
        path.write_text("## One\nUser: I do not recognise a charge.\nAgent: Confirm the last four.",
                        encoding="utf-8")
        complete, seen = _replying(_LAYOUT)
        conversations, _ = read_conversations(path, complete=complete)
        self.assertEqual(seen, [])
        self.assertTrue(conversations)


class TestTheSampleSentToTheCall(unittest.TestCase):
    def test_a_long_cell_is_cut_rather_than_sent_whole(self):
        """One row of a transcript export can be an entire conversation."""
        rows = [["id", "text"], ["C-1", "x" * 5000]]
        self.assertLess(len(conversation_migration.sample(rows)), 1000)

    def test_the_header_is_always_the_first_line(self):
        rows = [["id", "text"], ["C-1", "hello"]]
        self.assertTrue(conversation_migration.sample(rows).startswith("header: id | text"))

    def test_an_empty_file_is_described_rather_than_crashing(self):
        self.assertIn("no rows", conversation_migration.sample([]))
        self.assertIsNone(conversation_migration.migrate([], "empty.xlsx", lambda *a, **k: "{}"))


if __name__ == "__main__":
    unittest.main()
