"""The workbook this tool issues, coming back filled in.

This is the submission the tool actually asks for, and it was the shape the coverage reader
handled worst. There is no speaker column in a Variation_Log, because a turn is a *pair* of
columns rather than a row per utterance, so the generic reader fell through to
one-conversation-per-row: a hundred-row log became a hundred single-utterance "conversations"
whose text was whichever column happened to be widest, and a hundred model calls went out to work
out something the file states outright in its first column.

Three things are pinned here. The sheet is recognised. Its turns are reassembled into
conversations, one per scenario and variation, with both sides of each turn. And the scenario id
on the row is *honoured* rather than re-derived -- a submission returned whole costs no calls at
all, which is the difference between coverage being cheap and coverage being the most expensive
stage in the pipeline.

The last of those has a trap in it. Probes go out in the same template as the routes, so a team
can run a probe, say so on the row they were handed, and be perfectly well evidenced for it.
Reporting that as "matched no scenario" would describe a hole in their testing where the hole is
in this reader, and understating coverage is the one error this stage must not make quietly.
"""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scenario_generator.core.models import Capability, Decision, IntakeData, Persona, State, Tool
from scenario_generator.core.representation import build_report
from scenario_generator.ingest.conversations import read_conversations
from scenario_generator.io import read_space_metadata, write_data_template, write_space_metadata
from scenario_generator.llm.conversation_mapping import ConversationMapper
from scenario_generator.pipeline import build_scenario_space

_INTAKE = IntakeData(
    use_case={"Use case name": "Disputes", "Business objective": "Resolve disputes"},
    personas=[Persona("P1", "Cardmember", [], True)],
    capabilities=[Capability("CAP-01", "Identity", "Gating")],
    decisions=[Decision("DEC-01", "Identity check", "CAP-01", "", ["Pass", "Fail"])],
    states=[State("S-00", "Start", "The chat opens", ["DEC-01"], False),
            State("S-01", "DEC-01=Pass", "Verified", [], True, "Happy path"),
            State("S-02", "DEC-01=Fail", "Locked out", [], True, "Termination")],
    tools=[Tool("Identity service", "CAP-01", True)],
)


def _issued(directory: Path, with_probes: bool = False):
    """The two workbooks a real engagement has at the coverage stage."""
    scenarios = build_scenario_space(_INTAKE, with_probes=with_probes)
    for scenario in scenarios:
        scenario.description = f"A conversation exercising {scenario.id}."
        scenario.turn_plan = "1. Open the conversation.\n2. Answer the question."
    template = directory / "data_template.xlsx"
    metadata = directory / "scenario_space_metadata.xlsx"
    write_data_template(str(template), _INTAKE, scenarios)
    write_space_metadata(str(metadata), _INTAKE, scenarios)
    return template, metadata, scenarios


def _fill(template: Path, filled: Path, rows: int = None):
    """The log completed the way a team completes it: both halves of each turn."""
    book = load_workbook(template)
    log = book["Variation_Log"]
    header = [c.value for c in log[1]]
    user = header.index("User Input (Actual)") + 1
    agent = header.index("Agent Response") + 1
    last = min(log.max_row, (rows + 1) if rows else log.max_row)
    for row in range(2, last + 1):
        log.cell(row=row, column=user).value = "I do not recognise a charge on my statement."
        log.cell(row=row, column=agent).value = "Can you confirm the last four of your SSN?"
    book.save(filled)
    return filled


class TestTheSheetIsRecognised(unittest.TestCase):
    def setUp(self):
        self.directory = Path(tempfile.mkdtemp())
        self.template, self.metadata, self.scenarios = _issued(self.directory)
        self.filled = _fill(self.template, self.directory / "returned.xlsx")

    def test_it_says_it_read_the_template_rather_than_guessing_a_shape(self):
        _, how = read_conversations(self.filled)
        self.assertIn("data template returned filled in", how)

    def test_a_conversation_is_one_scenario_and_one_variation(self):
        conversations, _ = read_conversations(self.filled)
        book = load_workbook(self.template)
        log = book["Variation_Log"]
        header = [c.value for c in log[1]]
        pairs = {(r[0], r[1]) for r in log.iter_rows(min_row=2, values_only=True) if r[0]}
        self.assertEqual(len(conversations), len(pairs))

    def test_both_halves_of_a_turn_survive(self):
        """The generic reader took the widest column and labelled it User, which lost the agent's
        side of every exchange and mislabelled the half it kept."""
        conversations, _ = read_conversations(self.filled)
        first = conversations[0]
        self.assertEqual([t.speaker for t in first.turns[:2]], ["user", "agent"])
        self.assertIn("I do not recognise", first.turns[0].text)
        self.assertIn("last four", first.turns[1].text)

    def test_every_conversation_names_the_scenario_it_was_run_against(self):
        conversations, _ = read_conversations(self.filled)
        declared = {s.id for s in self.scenarios}
        for conversation in conversations:
            self.assertIn(conversation.scenario_id, declared)

    def test_a_template_nobody_filled_in_is_refused_with_a_reason(self):
        from scenario_generator.ingest.conversations import UnreadableConversations

        with self.assertRaises(UnreadableConversations) as caught:
            read_conversations(self.template)
        self.assertIn("no filled-in turns", str(caught.exception))

    def test_a_log_from_an_older_pack_still_reads(self):
        """The sheet was called Run_Log and its second column Run before variations were named."""
        book = load_workbook(self.filled)
        log = book["Variation_Log"]
        log.title = "Run_Log"
        log.cell(row=1, column=2).value = "Run"
        older = self.directory / "older.xlsx"
        book.save(older)

        conversations, how = read_conversations(older)
        self.assertIn("data template returned filled in", how)
        self.assertTrue(all(c.scenario_id for c in conversations))


class TestSomebodyElsesFormatStillReadsAsItDid(unittest.TestCase):
    """Recognising our own template must not take a file that is not it."""

    def test_a_turn_per_row_export_is_untouched(self):
        directory = Path(tempfile.mkdtemp())
        book = Workbook()
        sheet = book.active
        sheet.append(["Conversation ID", "Speaker", "Message"])
        for n in range(1, 3):
            sheet.append([f"C-{n}", "User", "I do not recognise a charge on my statement."])
            sheet.append([f"C-{n}", "Agent", "Can you confirm the last four of your SSN?"])
        path = directory / "theirs.xlsx"
        book.save(path)

        conversations, how = read_conversations(path)
        self.assertIn("one row per turn", how)
        self.assertEqual([c.scenario_id for c in conversations], ["", ""])


class TestAStatedScenarioIsNotRediscovered(unittest.TestCase):
    def setUp(self):
        self.directory = Path(tempfile.mkdtemp())
        self.template, self.metadata, self.scenarios = _issued(self.directory, with_probes=True)
        self.filled = _fill(self.template, self.directory / "returned.xlsx")
        self.conversations, _ = read_conversations(self.filled)
        self.calls = []

    def _map(self):
        def complete(system, user, **kwargs):
            self.calls.append(user)
            return "{}"

        return ConversationMapper(complete=complete).map(
            self.conversations, read_space_metadata(str(self.metadata)), _INTAKE,
            issued=read_space_metadata(str(self.metadata), functional_only=False))

    def test_a_submission_returned_whole_costs_no_calls(self):
        mappings = self._map()
        self.assertEqual(self.calls, [])
        self.assertEqual(len(mappings), len(self.conversations))

    def test_the_mapping_carries_the_id_the_row_stated(self):
        by_conversation = {m.conversation_id: m for m in self._map()}
        for conversation in self.conversations:
            self.assertEqual(by_conversation[conversation.id].scenario_id,
                             conversation.scenario_id)

    def test_the_mappings_come_back_in_submission_order(self):
        """They are written beside the conversations in the coverage workbook, so a list grouped
        by how each was worked out would read as a reordering of the model owner's file."""
        self.assertEqual([m.conversation_id for m in self._map()],
                         [c.id for c in self.conversations])

    def test_an_id_this_space_does_not_have_goes_to_the_call(self):
        """An id from a template issued off an older space is a real finding, not something to
        trust silently and not something to drop."""
        orphan = self.conversations[0]
        orphan.scenario_id = "SC-999"
        self._map()
        # One chunk call, and a solo retry because the stub's empty reply leaves it unanswered.
        # What matters is that exactly one conversation was ever asked about, and which one.
        self.assertTrue(self.calls, "the unknown id was trusted rather than checked")
        for sent in self.calls:
            self.assertIn(orphan.id, sent)
            for settled in self.conversations[1:]:
                self.assertNotIn(f'"{settled.id}"', sent)


class TestARunProbeIsEvidenceRatherThanAMiss(unittest.TestCase):
    def setUp(self):
        self.directory = Path(tempfile.mkdtemp())
        self.template, self.metadata, self.scenarios = _issued(self.directory, with_probes=True)
        self.filled = _fill(self.template, self.directory / "returned.xlsx")
        self.conversations, _ = read_conversations(self.filled)
        self.space = read_space_metadata(str(self.metadata))
        self.mappings = ConversationMapper(complete=lambda *a, **k: "{}").map(
            self.conversations, self.space, _INTAKE,
            issued=read_space_metadata(str(self.metadata), functional_only=False))

    def test_the_template_issues_probes_as_well_as_routes(self):
        """The premise of this whole class: a team is asked to run probes too."""
        self.assertTrue(any(c.scenario_id.startswith("NF-") for c in self.conversations))

    def test_a_probe_run_is_not_reported_as_matching_nothing(self):
        report = build_report(self.mappings, self.space)
        self.assertEqual(report.summary()["Matched no scenario"], 0)
        self.assertGreater(report.summary()["Ran a probe rather than a route"], 0)

    def test_probes_stay_out_of_what_coverage_measures(self):
        """A probe is a property of the agent rather than a route through it, so counting them in
        the denominator would move the figure without changing the evidence."""
        report = build_report(self.mappings, self.space)
        self.assertEqual(report.summary()["Scenarios in the space"], len(self.space))
        self.assertTrue(all(not s.scenario.id.startswith("NF-") for s in report.scenarios))

    def test_every_conversation_is_still_accounted_for_exactly_once(self):
        report = build_report(self.mappings, self.space)
        counts = report.summary()
        self.assertEqual(counts["Conversations read"], len(self.conversations))
        self.assertEqual(
            counts["Mapped to a scenario"] + counts["Matched no scenario"]
            + counts["Ran a probe rather than a route"], len(self.conversations))

    def test_a_conversation_that_genuinely_matched_nothing_still_reads_as_a_miss(self):
        """The distinction only helps if the other side of it survives."""
        for mapping in self.mappings:
            mapping.scenario_id = ""
        report = build_report(self.mappings, self.space)
        self.assertEqual(report.summary()["Matched no scenario"], len(self.mappings))
        self.assertNotIn("Ran a probe rather than a route", report.summary())


if __name__ == "__main__":
    unittest.main()
