"""Every scenario carries a short name, and the name survives every round trip.

A benchmark is three hundred rows in a spreadsheet and three hundred rows on a page, and until
now the only handle on any of them was the first sentence of its description. Those sentences all
begin the same way -- they are describing the same agent -- so the list was unscannable exactly
when it was large enough to need scanning.

The name is a handle rather than a summary: short, about the situation rather than the expected
behaviour, and different from its neighbours in the way the scenarios themselves differ. What is
pinned here is that one exists at every point in the pipeline, including before any model has
run, and that adding the column did not shift what an older registry reads back as.
"""
import tempfile
import unittest
from pathlib import Path

from scenario_generator.core.generation import fallback_name
from scenario_generator.core.models import (Capability, Decision, IntakeData, Persona, Scenario,
                                            State, Step, Tool, TurnMeta)
from scenario_generator.core.probes import build_probes
from scenario_generator.io import read_scenarios, write_registry
from scenario_generator.io.sheets import add_sheet, write_rows
from scenario_generator.llm.writer import MAX_NAME_CHARS, _clean_name
from scenario_generator.pipeline import build_scenarios

_INTAKE = IntakeData(
    use_case={"Use case name": "Disputes", "Business objective": "Resolve disputes"},
    personas=[Persona("P1", "Cardholder", [], True)],
    capabilities=[Capability("CAP-01", "Auth", "Gating")],
    decisions=[Decision("DEC-01", "Identity check", "CAP-01", "", ["Pass", "Fail"])],
    states=[State("S-00", "Start", "Start", ["DEC-01"], False),
            State("S-01", "Verified", "Happy path", [], True),
            State("S-02", "Locked out", "Termination", [], True)],
    tools=[Tool("Identity service", "CAP-01", True)],
)


class TestTheNameIsAHandle(unittest.TestCase):
    def test_a_trailing_full_stop_is_dropped(self):
        self.assertEqual(_clean_name("Locked out after two attempts."),
                         "Locked out after two attempts")

    def test_throat_clearing_is_stripped(self):
        """"Test that" is what the prompt asks for it not to say, and occasionally it says it."""
        for raw in ("Test that the account locks", "Verify that the account locks",
                    "Scenario where the account locks", "Check that the account locks"):
            self.assertEqual(_clean_name(raw), "The account locks", raw)

    def test_it_is_short_enough_to_sit_in_a_column(self):
        name = _clean_name("The cardmember " + "and again " * 30)
        self.assertLessEqual(len(name), MAX_NAME_CHARS)
        self.assertTrue(name.endswith("…"))

    def test_it_is_cut_at_a_word_rather_than_mid_word(self):
        """"the merchant wi…" reads as a rendering fault; "the merchant…" reads as a summary."""
        original = "Dispute filed against a merchant nobody at the bank has heard of before"
        kept = _clean_name(original).rstrip("…")
        self.assertTrue(original.startswith(kept))
        self.assertTrue(original[len(kept)] == " ", f"cut mid-word: {kept!r}")

    def test_it_is_one_line(self):
        self.assertEqual(_clean_name("Locked out\nafter two attempts"),
                         "Locked out after two attempts")

    def test_nothing_in_gives_nothing_out(self):
        self.assertEqual(_clean_name("   "), "")


class TestEveryScenarioHasOneBeforeAnyModelRuns(unittest.TestCase):
    """The walk is deterministic and so is the name it produces. A benchmark is scannable from
    the moment the graph is walked, not only after the writing stage has been paid for."""

    def test_a_walked_route_is_named_from_its_own_route(self):
        scenarios = build_scenarios(_INTAKE, with_probes=False)
        self.assertTrue(scenarios)
        for scenario in scenarios:
            self.assertTrue(scenario.name, f"{scenario.id} has no name")
            self.assertIn("Identity check", scenario.name)

    def test_a_probe_is_named_from_its_family(self):
        for probe in build_probes(_INTAKE):
            self.assertTrue(probe.name, f"{probe.id} has no name")

    def test_the_fallback_names_the_ending_not_the_beginning(self):
        """Two routes through the same first decision differ at the end, so the end has to be in
        the name or every scenario in a family reads identically."""
        meta = [TurnMeta(1, "DEC-01", "Identity check", "Pass", "", "S-01"),
                TurnMeta(2, "DEC-02", "Dispute eligible", "Too old", "", "S-05")]
        self.assertIn("Dispute eligible: Too old", fallback_name("Fallback", meta))


class TestTheNameSurvivesTheWorkbook(unittest.TestCase):
    def _round_trip(self, scenarios):
        path = Path(tempfile.mkdtemp()) / "registry.xlsx"
        write_registry(str(path), _INTAKE, scenarios)
        return {s.id: s for s in read_scenarios(str(path), _INTAKE)}

    def test_a_written_name_comes_back(self):
        scenarios = build_scenarios(_INTAKE, with_probes=False)
        scenarios[0].name = "Locked out after two failed checks"
        back = self._round_trip(scenarios)
        self.assertEqual(back[scenarios[0].id].name, "Locked out after two failed checks")

    def test_a_registry_written_before_names_existed_still_reads_correctly(self):
        """The Name column was inserted in the middle of the text sheet. Read by position, an
        older file would put its description into the name and its turn plan into the description
        -- silently, in the file that is the record of the whole benchmark."""
        from openpyxl import Workbook

        scenarios = build_scenarios(_INTAKE, with_probes=False)
        scenarios[0].description = "The cardmember disputes a charge."
        scenarios[0].turn_plan = "1. Open the conversation."

        path = Path(tempfile.mkdtemp()) / "old_registry.xlsx"
        write_registry(str(path), _INTAKE, scenarios)

        # Rewrite the text sheet in the shape it had before the Name column.
        from openpyxl import load_workbook
        book = load_workbook(path)
        del book["Scenario_Text"]
        sheet = add_sheet(book, "Scenario_Text", ["SC ID", "Description", "Turn Plan"],
                          [10, 60, 66])
        write_rows(sheet, [[s.id, s.description, s.turn_plan] for s in scenarios])
        book.save(path)

        back = {s.id: s for s in read_scenarios(str(path), _INTAKE)}[scenarios[0].id]
        self.assertEqual(back.description, "The cardmember disputes a charge.")
        self.assertEqual(back.turn_plan, "1. Open the conversation.")
        self.assertEqual(back.name, "")


class TestTheNameReachesThePageAndThePack(unittest.TestCase):
    def test_the_page_shows_the_name_as_the_handle(self):
        from scenario_generator.webapp.scenarios import to_row

        scenario = build_scenarios(_INTAKE, with_probes=False)[0]
        scenario.name = "Locked out after two failed checks"
        scenario.description = "The cardmember fails identity twice and the account locks."
        self.assertEqual(to_row(scenario)["title"], "Locked out after two failed checks")

    def test_the_challenge_pack_carries_it_beside_the_id(self):
        from scenario_generator.io import write_challenge_pack
        from scenario_generator.io.sheets import open_for_reading, read_table

        scenarios = build_scenarios(_INTAKE, with_probes=False)
        scenarios[0].name = "Locked out after two failed checks"
        path = Path(tempfile.mkdtemp()) / "pack.xlsx"
        write_challenge_pack(str(path), _INTAKE, scenarios)

        with open_for_reading(str(path), "a challenge pack") as book:
            header, rows = read_table(book["Scenarios"])
        self.assertEqual(header[:3], ["SC ID", "Name", "Description"])
        self.assertIn("Locked out after two failed checks", [row[1] for row in rows])

    def test_the_pack_still_carries_no_expected_outcome(self):
        """Adding a column to the issued pack is exactly where an answer key could slip in."""
        from scenario_generator.io import write_challenge_pack
        from scenario_generator.io.sheets import open_for_reading, read_table

        scenarios = build_scenarios(_INTAKE, with_probes=False)
        path = Path(tempfile.mkdtemp()) / "pack.xlsx"
        write_challenge_pack(str(path), _INTAKE, scenarios)

        with open_for_reading(str(path), "a challenge pack") as book:
            text = " ".join(str(cell) for sheet in book.worksheets
                            for row in read_table(sheet)[1] for cell in row)
        for terminal in ("Verified", "Locked out"):
            self.assertNotIn(terminal, text, f"the pack names the ending state {terminal!r}")


if __name__ == "__main__":
    unittest.main()
