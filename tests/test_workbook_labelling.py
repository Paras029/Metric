"""The sensitivity label every workbook this tool writes carries.

A workbook written by a library arrives with no classification on it. Where labelling is mandatory
that is not a neutral state: the file is unlabelled, so the organisation's default applies, and
where that default is the most restrictive one the result is a registry nobody can open. The whole
point of these tests is that "no label" and "a permissive label" are opposite outcomes, and only
one of them is what leaving a field blank looks like.

Two rules carry the behaviour.

*Every workbook this tool writes gets the configured label.* Not the registry alone -- the template,
the drafted intake, the challenge pack, the coverage report and the graph file all leave here, and
one of them being unopenable is as bad as all of them.

*A workbook that already carries a label is never relabelled.* Re-saving an intake somebody
uploaded must not silently reclassify their document, whichever direction that would move it.
"""
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.packaging.custom import StringProperty

from scenario_generator.core.models import IntakeData, Persona
from scenario_generator.io import labelling
from scenario_generator.llm import config

_LABEL = "0e5cf1c2-1a2b-4c3d-9e4f-a5b6c7d8e9f0"
_INTAKE = IntakeData(use_case={"Use case name": "Claims"},
                     personas=[Persona("P1", "Default", ["wants to file a claim"], True)],
                     capabilities=[], decisions=[], states=[], tools=[])


def _labels_on(path) -> dict:
    """Every sensitivity property on a workbook, read back off the file itself."""
    return {str(p.name): str(p.value) for p in load_workbook(path).custom_doc_props.props
            if str(p.name).startswith(labelling.PREFIX)}


class _Configured(unittest.TestCase):
    def setUp(self):
        self.work = Path(tempfile.mkdtemp())
        self.tuning = self.work / "tuning.yml"
        self._environment = mock.patch.dict(os.environ, {"TUNING_PATH": str(self.tuning)})
        self._environment.start()
        self.addCleanup(self._environment.stop)
        self.addCleanup(config.reload_tuning)
        self.addCleanup(labelling.forget_warnings)
        self.configure({"enabled": True, "label_id": _LABEL, "name": "AXP Internal",
                        "site_id": "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"})

    def configure(self, sensitivity: dict) -> None:
        self.tuning.write_text(yaml.safe_dump({"sensitivity": sensitivity}), encoding="utf-8")
        config.reload_tuning()
        labelling.forget_warnings()


class TestEveryWorkbookLeavesLabelled(_Configured):
    def _written(self, name: str) -> Path:
        """One of each kind of workbook this tool produces."""
        from scenario_generator.core.intake import write_template
        from scenario_generator.ingest import DraftedIntake, write_drafted_intake
        from scenario_generator.ingest.drafting import _validate
        from scenario_generator.io import (write_challenge_pack, write_coverage_report,
                                           write_registry, write_scenario_graph)
        from scenario_generator.core.representation import CoverageReport

        path = self.work / f"{name}.xlsx"
        if name == "template":
            write_template(str(path))
        elif name == "registry":
            write_registry(str(path), _INTAKE, [])
        elif name == "challenge_pack":
            write_challenge_pack(str(path), _INTAKE, [])
        elif name == "scenario_graph":
            write_scenario_graph(str(path), _INTAKE, [])
        elif name == "coverage":
            write_coverage_report(str(path), CoverageReport(), [])
        elif name == "drafted_intake":
            write_drafted_intake(path, DraftedIntake(_validate(
                {"use_case": {"name": "Claims"}, "personas": [], "capabilities": [],
                 "decisions": [], "states": [], "tools": []})))
        return path

    def test_the_registry_is_labelled(self):
        """The one that was reported: a registry nobody could open."""
        labels = _labels_on(self._written("registry"))
        self.assertEqual(labels[f"{labelling.PREFIX}{_LABEL}_Name"], "AXP Internal")
        self.assertEqual(labels[f"{labelling.PREFIX}{_LABEL}_Enabled"], "true")

    def test_every_other_workbook_is_labelled_too(self):
        """One unopenable output is as bad as all of them, so none of them may be missed."""
        for name in ("template", "challenge_pack", "scenario_graph", "coverage",
                     "drafted_intake"):
            with self.subTest(workbook=name):
                labels = _labels_on(self._written(name))
                self.assertEqual(labels.get(f"{labelling.PREFIX}{_LABEL}_Name"), "AXP Internal")

    def test_the_label_is_in_the_file_where_office_looks_for_it(self):
        """Custom document properties, in the package part Office reads them from."""
        with zipfile.ZipFile(self._written("registry")) as archive:
            self.assertIn("docProps/custom.xml", archive.namelist())
            self.assertIn(_LABEL, archive.read("docProps/custom.xml").decode())

    def test_the_workbook_still_reads_back_as_an_intake(self):
        """A label must not cost the file its own job."""
        from scenario_generator.core.intake import read_intake
        from scenario_generator.ingest import DraftedIntake, write_drafted_intake
        from scenario_generator.ingest.drafting import _validate

        path = self.work / "round_trip.xlsx"
        write_drafted_intake(path, DraftedIntake(_validate(
            {"use_case": {"name": "Claims"},
             "personas": [{"id": "P1", "name": "Default", "applies_to": "wants to file",
                           "is_default": True}],
             "capabilities": [], "decisions": [], "states": [], "tools": []})))
        self.assertEqual(read_intake(str(path)).name, "Claims")
        self.assertTrue(_labels_on(path))

    def test_the_method_says_a_tool_applied_it_rather_than_a_person(self):
        """"Privileged" means somebody chose the label deliberately. Claiming that would be
        untrue, and it is the field a reviewer would check."""
        labels = _labels_on(self._written("registry"))
        self.assertEqual(labels[f"{labelling.PREFIX}{_LABEL}_Method"], "Standard")


class TestALabelAlreadyThereIsLeftAlone(_Configured):
    def test_a_workbook_carrying_its_own_label_is_not_relabelled(self):
        """Re-saving an intake somebody uploaded must not silently reclassify their document."""
        theirs = "99999999-8888-7777-6666-555555555555"
        book = Workbook()
        for field, value in (("Enabled", "true"), ("Name", "AXP Confidential")):
            book.custom_doc_props.append(
                StringProperty(name=f"{labelling.PREFIX}{theirs}_{field}", value=value))

        self.assertFalse(labelling.apply(book))
        names = {p.name for p in book.custom_doc_props.props}
        self.assertIn(f"{labelling.PREFIX}{theirs}_Name", names)
        self.assertNotIn(f"{labelling.PREFIX}{_LABEL}_Name", names)

    def test_an_unlabelled_workbook_is_stamped(self):
        book = Workbook()
        self.assertTrue(labelling.apply(book))
        self.assertEqual(labelling.existing_label(book), _LABEL)

    def test_an_uploaded_intake_keeps_its_label_through_an_edit(self):
        """set_decision_scope writes to somebody's own workbook in place. The classification on it
        is theirs and has to survive that."""
        from scenario_generator.core.intake import set_decision_scope, write_template

        # Built as *their* file: written with labelling off, then labelled the way an upload from
        # somebody else would already be. A file carrying two labels is not a real case -- nothing
        # here ever adds a second -- so the fixture must not manufacture one.
        path = self.work / "theirs.xlsx"
        self.configure({"enabled": False})
        write_template(str(path))
        self.configure({"enabled": True, "label_id": _LABEL, "name": "AXP Internal"})

        book = load_workbook(path)
        book["L3 Decisions"].append(["DEC-01", "Check", "", "", "Pass / Fail", "User", 1, "", "No"])
        for field, value in (("Enabled", "true"), ("Name", "AXP Confidential")):
            book.custom_doc_props.append(StringProperty(
                name=f"{labelling.PREFIX}deadbeef-0000-1111-2222-333333333333_{field}",
                value=value))
        book.save(path)

        set_decision_scope(str(path), "DEC-01", True)

        labels = _labels_on(path)
        self.assertEqual(labels.get(f"{labelling.PREFIX}deadbeef-0000-1111-2222-333333333333_Name"),
                         "AXP Confidential")
        self.assertNotIn(f"{labelling.PREFIX}{_LABEL}_Name", labels)


class TestSupplyingTheLabelWithoutKnowingItsGuid(_Configured):
    """A label id is a GUID minted in the tenant. Nothing guesses it, so there has to be a way to
    supply it that does not require going and finding it."""

    def _reference(self, name: str = "AXP Internal") -> Path:
        path = self.work / "already_labelled.xlsx"
        book = Workbook()
        for field, value in (("Enabled", "true"), ("Name", name), ("Method", "Standard"),
                             ("SiteId", "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee")):
            book.custom_doc_props.append(
                StringProperty(name=f"{labelling.PREFIX}{_LABEL}_{field}", value=value))
        book.save(path)
        return path

    def test_the_label_is_read_off_a_workbook_that_already_has_the_right_one(self):
        reference = self._reference()
        self.configure({"enabled": True, "copy_from": str(reference)})

        from scenario_generator.io import write_registry

        registry = self.work / "registry.xlsx"
        write_registry(str(registry), _INTAKE, [])
        self.assertEqual(_labels_on(registry)[f"{labelling.PREFIX}{_LABEL}_Name"], "AXP Internal")

    def test_copying_wins_over_an_id_written_out_by_hand(self):
        """Two ways to say the same thing; the one read off a real file is the one to trust."""
        self.configure({"enabled": True, "copy_from": str(self._reference("From the file")),
                        "label_id": "11111111-1111-1111-1111-111111111111", "name": "Typed"})
        properties = labelling.label_properties()
        self.assertEqual(properties[f"{labelling.PREFIX}{_LABEL}_Name"], "From the file")

    def test_a_reference_file_with_no_label_is_reported_rather_than_guessed_at(self):
        bare = self.work / "unlabelled.xlsx"
        Workbook().save(bare)
        self.configure({"enabled": True, "copy_from": str(bare)})
        with self.assertLogs("scenario_generator.io.labelling", level="WARNING"):
            self.assertEqual(labelling.label_properties(), {})


class TestWhenNoLabelIsConfigured(_Configured):
    def test_nothing_is_invented(self):
        """A made-up GUID names a label that does not exist in the tenant, which fails differently
        and worse than no label at all."""
        self.configure({"enabled": True, "label_id": "", "copy_from": ""})
        with self.assertLogs("scenario_generator.io.labelling", level="WARNING"):
            self.assertEqual(labelling.label_properties(), {})

    def test_the_run_says_so_rather_than_writing_a_file_nobody_can_open(self):
        self.configure({"enabled": True, "label_id": "", "copy_from": ""})
        with self.assertLogs("scenario_generator.io.labelling", level="WARNING") as logged:
            labelling.label_properties()
        message = "\n".join(logged.output)
        self.assertIn("most restrictive", message)
        self.assertIn("copy_from", message)

    def test_the_warning_is_said_once_rather_than_per_workbook(self):
        self.configure({"enabled": True, "label_id": "", "copy_from": ""})
        with self.assertLogs("scenario_generator.io.labelling", level="WARNING") as logged:
            for _ in range(5):
                labelling.label_properties()
        self.assertEqual(len(logged.output), 1)

    def test_writing_still_works(self):
        """An unlabelled workbook is a problem for the reader, not a reason to fail the run."""
        self.configure({"enabled": True, "label_id": "", "copy_from": ""})
        from scenario_generator.io import write_registry

        registry = self.work / "registry.xlsx"
        write_registry(str(registry), _INTAKE, [])
        self.assertTrue(registry.exists())
        self.assertEqual(_labels_on(registry), {})

    def test_labelling_can_be_turned_off_outright(self):
        self.configure({"enabled": False, "label_id": _LABEL, "name": "AXP Internal"})
        book = Workbook()
        self.assertFalse(labelling.apply(book))
        self.assertIn("off", labelling.describe())


class TestTheSettingIsReachableTheUsualWays(_Configured):
    def test_the_environment_wins_over_the_file(self):
        with mock.patch.dict(os.environ, {"SENSITIVITY_NAME": "AXP Internal - Override"}):
            self.assertEqual(config.SENSITIVITY_NAME, "AXP Internal - Override")

    def test_the_shipped_tuning_file_carries_the_section(self):
        shipped = yaml.safe_load(
            (Path(__file__).resolve().parent.parent / "tuning.yml").read_text(encoding="utf-8"))
        self.assertIn("sensitivity", shipped)
        for key in ("enabled", "copy_from", "label_id", "name", "method"):
            self.assertIn(key, shipped["sensitivity"])


if __name__ == "__main__":
    unittest.main()
