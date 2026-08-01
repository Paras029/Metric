"""Reading a submitted pack into answers about the agent.

The documents are read whole rather than passage by passage, so the tests that matter are about
what that buys and what it must not cost: a fact spread over three sections assembled into one
answer, a citation checked against the source, and a small, predictable number of model calls.

`test_a_pack_is_read_in_a_handful_of_calls` is the one guarding the call count. A per-passage
reading costs roughly thirty calls for a sixty-page pack, and every one of those is a wait and a
chance to fail.
"""
import json
import tempfile
import unittest
from pathlib import Path

from scenario_generator.ingest import (IngestionFailed, build_context_document, build_corpus,
                                       extract_documents, open_questions, record_from_json,
                                       record_to_json)

# One process described across three separate sections, the way real documentation does it.
_SCATTERED = """# Handling a request

The assistant first confirms who it is speaking to before anything else happens.

# Verification

Where confirmation fails the assistant may try again. After the third unsuccessful attempt the
session is locked.

# Escalation

A locked session is passed to a human agent, who takes over the conversation.
"""

_ANSWER = {
    "answer": "The agent confirms identity, retries up to three times, then locks and escalates.",
    "points": ["Identity is confirmed first.", "Three attempts are allowed.",
               "A locked session goes to a human agent."],
    "unknowns": ["Which identifiers are accepted for confirmation?"],
    "evidence": [{"quote": "The assistant first confirms who it is speaking to",
                  "document": "notes.md", "locator": "under “Handling a request”"},
                 {"quote": "A locked session is passed to a human agent",
                  "document": "notes.md", "locator": "under “Escalation”"}],
    "confidence": "High",
}


def _write(name, body):
    path = Path(tempfile.mkdtemp()) / name
    path.write_text(body, encoding="utf-8")
    return path


def _stub(answer=None, facets=("decisions",), resolved=None, calls=None):
    """A completion that answers whichever prompt it is handed."""
    answer = answer if answer is not None else _ANSWER

    def complete(system, user, **kwargs):
        if calls is not None:
            calls.append(user)
        if "THE OUTSTANDING QUESTIONS" in user:
            return json.dumps({"resolved": resolved or []})
        # A reading call names the questions it wants answered; reply only to those.
        return json.dumps({facet: answer for facet in facets if f"- {facet}:" in user})
    return complete


class TestCorpus(unittest.TestCase):
    def test_every_document_is_named_and_its_locators_kept(self):
        corpus = build_corpus([_write("notes.md", _SCATTERED)])
        self.assertIn("=== DOCUMENT: notes.md ===", corpus.text)
        self.assertIn("[under “Escalation”]", corpus.text)

    def test_an_unreadable_file_is_recorded_and_the_rest_still_read(self):
        corpus = build_corpus([_write("empty.md", "   "), _write("notes.md", _SCATTERED)])
        kinds = {d.name: d.kind for d in corpus.documents}
        self.assertEqual(kinds["empty.md"], "unreadable")
        self.assertIn("notes.md", corpus.text)

    def test_images_are_set_aside_for_the_vision_pass(self):
        image = Path(tempfile.mkdtemp()) / "flow.png"
        image.write_bytes(b"not really a png")
        corpus = build_corpus([image])
        self.assertEqual([p.name for p in corpus.diagrams], ["flow.png"])


class TestReading(unittest.TestCase):
    def test_a_pack_is_read_in_a_handful_of_calls(self):
        """Three reading calls and two resolution sweeps, however long the document is.

        The number matters because reading a passage at a time would make it roughly thirty for
        a sixty-page pack. It is fixed by how many groups of questions there are and how many
        times what is still open is put back to the documents -- not by the length of the pack.
        """
        calls = []
        extract_documents([_write("notes.md", _SCATTERED)], complete=_stub(calls=calls))
        self.assertLessEqual(len(calls), 5, f"expected a handful of calls, made {len(calls)}")

    def test_scattered_facts_are_assembled_into_one_answer(self):
        record = extract_documents([_write("notes.md", _SCATTERED)], complete=_stub())
        answer = record.answer_for("decisions")
        self.assertTrue(answer.is_answered)
        self.assertIn("three times", answer.answer)
        self.assertEqual(len(answer.points), 3)

    def test_a_citation_found_in_the_documents_becomes_a_traceable_claim(self):
        record = extract_documents([_write("notes.md", _SCATTERED)], complete=_stub())
        self.assertTrue(record.claims)
        self.assertTrue(record.answer_for("decisions").sources)
        self.assertIn(record.answer_for("decisions").sources[0], record.claims_by_id())

    def test_an_invented_citation_is_dropped(self):
        answer = dict(_ANSWER, evidence=[{
            "quote": "Cardmembers may raise a dispute within ninety days of the posting date.",
            "document": "notes.md", "locator": "x"}])
        record = extract_documents([_write("notes.md", _SCATTERED)],
                                   complete=_stub(answer=answer))
        self.assertEqual(record.claims, [])

    def test_an_answer_with_no_surviving_citation_is_flagged_rather_than_trusted(self):
        answer = dict(_ANSWER, evidence=[{"quote": "Nothing like this appears anywhere at all.",
                                          "document": "notes.md", "locator": "x"}])
        record = extract_documents([_write("notes.md", _SCATTERED)],
                                   complete=_stub(answer=answer))
        self.assertTrue(any("None of the quotes" in u
                            for u in record.answer_for("decisions").unknowns))

    def test_a_question_nothing_answered_reports_itself_unanswered(self):
        record = extract_documents([_write("notes.md", _SCATTERED)], complete=_stub())
        self.assertIn("tools", record.empty_facets())

    def test_a_run_that_mostly_failed_is_abandoned_rather_than_written(self):
        def broken(system, user, **kwargs):
            raise RuntimeError("401 Unauthorized")
        with self.assertRaises(IngestionFailed):
            extract_documents([_write("notes.md", _SCATTERED)], complete=broken)

    def test_a_pack_where_nothing_could_be_read_fails_before_calling_anything(self):
        with self.assertRaises(IngestionFailed):
            extract_documents([_write("empty.md", "  ")], complete=_stub())


class TestResolutionSweep(unittest.TestCase):
    """Every question that survives to the model owner costs days, so they are asked twice."""

    def test_a_second_look_can_settle_what_the_first_reading_left_open(self):
        resolved = [{"question": "Which identifiers are accepted for confirmation?",
                     "status": "answered",
                     "answer": "An account number and a postcode.",
                     "evidence": [], "still_open": ""}]
        record = extract_documents([_write("notes.md", _SCATTERED)],
                                   complete=_stub(resolved=resolved))
        answer = record.answer_for("decisions")
        self.assertIn("An account number and a postcode.", answer.points)
        self.assertNotIn("Which identifiers are accepted for confirmation?", answer.unknowns)

    def test_a_question_the_documents_do_not_settle_stays_open(self):
        resolved = [{"question": "Which identifiers are accepted for confirmation?",
                     "status": "unanswered", "answer": "", "evidence": [], "still_open": ""}]
        record = extract_documents([_write("notes.md", _SCATTERED)],
                                   complete=_stub(resolved=resolved))
        self.assertIn("Which identifiers are accepted for confirmation?",
                      record.answer_for("decisions").unknowns)

    def test_the_sweep_can_be_switched_off(self):
        calls = []
        extract_documents([_write("notes.md", _SCATTERED)], resolve=False,
                          complete=_stub(calls=calls))
        self.assertFalse(any("THE OUTSTANDING QUESTIONS" in c for c in calls))


class TestDiagrams(unittest.TestCase):
    def _png(self):
        import base64
        path = Path(tempfile.mkdtemp()) / "flow.png"
        path.write_bytes(base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmM"
            "IQAAAABJRU5ErkJggg=="))
        return path

    # A whole, self-consistent workflow: one decision with both outcomes leading somewhere, and
    # a start. Complete on purpose -- these tests are about what happens to the reading, and a
    # structure with holes in it would drag the repair pass into every one of them.
    _WHOLE = {
        "capabilities": [{"id": "CAP-01", "name": "Authentication", "type": "Gating"}],
        "decisions": [{"id": "DEC-01", "name": "PIN check", "outcomes": ["Pass", "Fail"],
                       "capability_id": "CAP-01", "input_source": "User", "max_attempts": 1}],
        "states": [
            {"id": "S-00", "reached_via": "Start", "description": "Call opens",
             "next_decisions": ["DEC-01"], "is_terminal": False},
            {"id": "S-01", "reached_via": "DEC-01=Pass", "description": "Verified",
             "next_decisions": [], "is_terminal": True, "outcome_type": "Happy path"},
            {"id": "S-02", "reached_via": "DEC-01=Fail", "description": "Locked",
             "next_decisions": [], "is_terminal": True, "outcome_type": "Termination"}],
    }

    def _boxes(self, system, user, images, **kwargs):
        """A per-image reading in the shape the read prompt asks for."""
        return json.dumps({
            "nodes": [{"ref": "n1", "label": "PIN check", "kind": "decision"}],
            "edges": [{"from": "n1", "to": "n2", "label": "Pass"}],
            "continues_offpage": [], "unreadable": [],
            "counts": {"boxes": 1, "arrows": 1}})

    def _synthesizing(self, observations, base=None, structure=None):
        """A completion answering the ordinary reading prompts through ``base`` (defaulting to
        ``_stub()``), and the diagram-synthesis prompt with a graph plus the given observations."""
        base = base or _stub()
        graph = self._WHOLE if structure is None else structure

        def complete(system, user, **kwargs):
            if "THE READINGS" in user:
                return json.dumps({**graph, "observations": observations})
            return base(system, user, **kwargs)
        return complete

    def test_a_diagram_is_read_and_marked_for_confirmation(self):
        def describe_one(system, user, images, **kwargs):
            assert images and images[0][0] == "image/png"
            return self._boxes(system, user, images, **kwargs)

        complete = self._synthesizing([{
            "facet": "decisions", "statement": "Auth check branches to pass or fail.",
            "quote": "Auth check", "locator": "top left"}])

        record = extract_documents([_write("notes.md", _SCATTERED), self._png()],
                                   complete=complete, describe_images=describe_one)
        diagram_claims = [c for c in record.claims if c.source.kind == "image"]
        self.assertTrue(diagram_claims)
        self.assertTrue(diagram_claims[0].needs_confirmation)

    def test_what_a_diagram_establishes_reaches_the_context_document(self):
        """The whole point of reading a diagram.

        The context document renders only the claims an *answer* names as its sources, and the
        intake is drafted from the context document alone -- so an observation stored on the
        record but named by no answer is read by nothing. A workflow diagram is often the only
        place a branch is written down, and a pack whose structure lives entirely in its images
        must not produce a context file reporting that structure as "not covered by the submitted
        documents".
        """
        describe_one = self._boxes

        complete = self._synthesizing([{
            "facet": "decisions", "statement": "PIN check branches to Pass or Fail.",
            "quote": "PIN check", "locator": "image 1"}])

        record = extract_documents([_write("notes.md", _SCATTERED), self._png()],
                                   complete=complete, describe_images=describe_one)

        answer = record.answer_for("decisions")
        self.assertIn("PIN check branches to Pass or Fail.", answer.points)
        self.assertTrue(answer.is_answered)
        self.assertNotIn("decisions", record.empty_facets())

        # Traceable as well as present: the claim carries an id and the answer names it.
        claim = next(c for c in record.claims if c.source.kind == "image")
        self.assertTrue(claim.id)
        self.assertIn(claim.id, answer.sources)

        context = build_context_document(record, "Test agent")
        self.assertIn("PIN check branches to Pass or Fail.", context)

    def test_a_facet_only_a_diagram_answered_stops_being_asked_as_a_gap(self):
        """The placeholder question means 'nothing addressed this at all', which stops being
        true once a diagram has. Leaving it would ask the model owner for what was already
        sent."""
        describe_one = self._boxes

        complete = self._synthesizing([{
            "facet": "states", "statement": "A Locked state ends the call.",
            "quote": "Locked", "locator": "image 1"}])

        record = extract_documents([_write("notes.md", _SCATTERED), self._png()],
                                   complete=complete, describe_images=describe_one)

        generic = "What positions can an interaction be in, and which of them end it?"
        self.assertNotIn(generic, record.answer_for("states").unknowns)

    def test_diagram_confirmations_are_grouped_by_facet_not_one_per_observation(self):
        """A single diagram can produce dozens of observations; each becoming its own question
        would flood the open-questions list with rows that all ask the same thing."""
        describe_one = self._boxes

        complete = self._synthesizing([
            {"facet": "decisions", "statement": "Auth check branches to pass or fail.",
             "quote": "Auth check", "locator": "top left"},
            {"facet": "decisions", "statement": "Escalation follows three failures.",
             "quote": "Escalate", "locator": "bottom right"},
            {"facet": "states", "statement": "A locked state ends the call.",
             "quote": "Locked", "locator": "bottom"},
        ])

        record = extract_documents([_write("notes.md", _SCATTERED), self._png()],
                                   complete=complete, describe_images=describe_one)
        confirmations = [q for q in open_questions(record) if q["kind"] == "confirm"]

        self.assertEqual(len(confirmations), 2)               # one per facet, not one per claim
        decisions = next(q for q in confirmations if q["facet"] == "decisions")
        self.assertIn("Auth check branches to pass or fail.", decisions["detail"])
        self.assertIn("Escalation follows three failures.", decisions["detail"])
        self.assertIn("2 statement", decisions["question"])

    def test_several_images_are_read_on_their_own_then_put_together(self):
        """Each image is described without seeing the others; the synthesis call is given every
        description together, and is the only one that has to make sense of all of them at once."""
        reads = []

        def describe_one(system, user, images, **kwargs):
            reads.append(images[0])
            return json.dumps({
                "nodes": [{"ref": "n1", "label": f"Box in image {len(reads)}",
                           "kind": "decision"}],
                "edges": [], "continues_offpage": [], "unreadable": [],
                "counts": {"boxes": 1, "arrows": 0}})

        synth_calls = []

        def complete(system, user, **kwargs):
            if "THE READINGS" in user:
                synth_calls.append(user)
                return json.dumps({**self._WHOLE, "observations": [{
                    "facet": "decisions", "statement": "Combined across both images.",
                    "quote": "Auth check", "locator": "image 2"}]})
            return _stub()(system, user, **kwargs)

        second = Path(tempfile.mkdtemp()) / "flow2.png"
        second.write_bytes(self._png().read_bytes())

        record = extract_documents([_write("notes.md", _SCATTERED), self._png(), second],
                                   complete=complete, describe_images=describe_one)

        self.assertEqual(len(reads), 2)                    # each image read on its own
        self.assertIn("Box in image 1", synth_calls[0])
        self.assertIn("Box in image 2", synth_calls[0])
        # Node references are namespaced per image, so the join can tell them apart.
        self.assertIn("i1.n1", synth_calls[0])
        self.assertIn("i2.n1", synth_calls[0])
        diagram_claims = [c for c in record.claims if c.source.kind == "image"]
        self.assertTrue(diagram_claims)

    def test_an_unreadable_image_is_dropped_and_the_rest_still_go_through(self):
        """One image fails on its own; the other is still read and still reaches synthesis."""
        def describe_one(system, user, images, **kwargs):
            if any(image[1] == b"broken" for image in images):
                raise RuntimeError("vision not enabled")
            return self._boxes(system, user, images, **kwargs)

        complete = self._synthesizing([{
            "facet": "decisions", "statement": "From the one readable diagram.",
            "quote": "Auth check", "locator": "top left"}])

        broken_image = Path(tempfile.mkdtemp()) / "broken.png"
        broken_image.write_bytes(b"broken")

        record = extract_documents(
            [_write("notes.md", _SCATTERED), self._png(), broken_image],
            complete=complete, describe_images=describe_one)

        self.assertTrue(any(c.source.kind == "image" for c in record.claims))

    def test_an_unreadable_diagram_asks_for_a_description_and_the_pack_still_reads(self):
        def broken(system, user, images, **kwargs):
            raise RuntimeError("vision not enabled")

        record = extract_documents([_write("notes.md", _SCATTERED), self._png()],
                                   complete=_stub(), describe_images=broken)
        diagram = next(d for d in record.documents if d.name == "flow.png")
        self.assertEqual(diagram.kind, "unreadable")
        self.assertIn("written description", diagram.note)
        self.assertTrue(record.answer_for("decisions").is_answered)

    def test_a_synthesis_failure_still_leaves_the_rest_of_the_pack_readable(self):
        """Every image was read individually, but the pass that puts them together failed."""
        describe_one = self._boxes

        def complete(system, user, **kwargs):
            if "THE READINGS" in user:
                raise RuntimeError("gateway down")
            return _stub()(system, user, **kwargs)

        record = extract_documents([_write("notes.md", _SCATTERED), self._png()],
                                   complete=complete, describe_images=describe_one)
        diagram = next(d for d in record.documents if d.name == "flow.png")
        self.assertEqual(diagram.kind, "unreadable")
        self.assertTrue(record.answer_for("decisions").is_answered)


class TestContextDocument(unittest.TestCase):
    def _record(self):
        return extract_documents([_write("notes.md", _SCATTERED)], complete=_stub())

    def test_the_assembled_answer_leads_the_document(self):
        text = build_context_document(self._record(), "Disputes assistant")
        self.assertIn("retries up to three times", text)
        self.assertIn("Where it branches", text)

    def test_supporting_quotes_are_traceable(self):
        text = build_context_document(self._record())
        self.assertIn("notes.md", text)

    def test_what_is_unanswered_is_stated_rather_than_omitted(self):
        self.assertIn("Not covered by the submitted documents",
                      build_context_document(self._record()))

    def test_open_questions_are_listed_without_duplication(self):
        questions = [q["question"] for q in open_questions(self._record())]
        self.assertEqual(len(questions), len(set(questions)))

    def test_the_record_round_trips_through_json(self):
        record = self._record()
        target = Path(tempfile.mkdtemp()) / "evidence.json"
        record_to_json(record, target)
        restored = record_from_json(target)
        self.assertEqual(restored.answer_for("decisions").answer,
                         record.answer_for("decisions").answer)


if __name__ == "__main__":
    unittest.main()


class TestCoverageAnnotation(unittest.TestCase):
    """Coverage is recorded against a scenario and shown to a person. It removes nothing."""

    def test_the_annotation_round_trips_and_stays_out_of_the_pack(self):
        import sys
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        from openpyxl import load_workbook
        from test_stage_runners import _intake_workbook
        from scenario_generator.core.intake import read_intake
        from scenario_generator.io import read_scenarios, write_challenge_pack, write_registry
        from scenario_generator.pipeline import build_scenarios

        directory = Path(tempfile.mkdtemp())
        intake = read_intake(str(_intake_workbook(directory)))
        scenarios = build_scenarios(intake, with_probes=True)
        scenarios[0].owner_coverage = "Covered"
        scenarios[0].owner_coverage_note = "Their TC-001"

        write_registry(str(directory / "registry.xlsx"), intake, scenarios)
        restored = read_scenarios(str(directory / "registry.xlsx"), intake)
        self.assertEqual(restored[0].owner_coverage, "Covered")
        self.assertEqual(restored[0].owner_coverage_note, "Their TC-001")

        # The model owner must not learn which scenarios are already covered -- that would say
        # which ones the validator considers already answered.
        write_challenge_pack(str(directory / "pack.xlsx"), intake, scenarios)
        book = load_workbook(directory / "pack.xlsx")
        values = [str(v) for name in book.sheetnames
                  for row in book[name].iter_rows(values_only=True) for v in row if v]
        self.assertFalse(any("Covered" in v for v in values))

    def test_annotating_records_the_count_and_removes_nothing(self):
        import sys
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        from test_stage_runners import _intake_workbook
        from scenario_generator.core.intake import read_intake
        from scenario_generator.core.representation import build_report
        from scenario_generator.io import read_registry, read_scenarios, write_registry
        from scenario_generator.llm.conversation_mapping import Mapping
        from scenario_generator.pipeline import annotate_coverage, build_scenarios

        directory = Path(tempfile.mkdtemp())
        registry = str(directory / "registry.xlsx")
        intake = read_intake(str(_intake_workbook(directory)))
        scenarios = build_scenarios(intake, with_probes=True)
        write_registry(registry, intake, scenarios)

        benchmark = read_registry(registry)
        report = build_report([Mapping("C1", benchmark[0].id, "high"),
                               Mapping("C2", benchmark[0].id, "low")], benchmark)

        self.assertEqual(annotate_coverage(registry, intake, report), 1)
        restored = read_scenarios(registry, intake)
        self.assertEqual(len(restored), len(scenarios))
        annotated = next(s for s in restored if s.id == benchmark[0].id)
        self.assertEqual(annotated.owner_coverage, "2 conversations")
        self.assertIn("C1", annotated.owner_coverage_note)
