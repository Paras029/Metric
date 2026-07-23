"""The challenge pack is the one artifact that leaves MRMG. These tests lock its shape and,
most importantly, that it carries no ground truth.
"""
import os
import tempfile
import unittest

from openpyxl import load_workbook

from scenario_generator.core.generation import recommended_turns, required_runs
from scenario_generator.core.models import (Decision, IntakeData, Persona, State, Tool)
from scenario_generator.core.probes import build_probes
from scenario_generator.io import write_challenge_pack

_INTAKE = IntakeData(
    use_case={"Use case name": "Test", "Business objective": "Objective"},
    personas=[Persona("P1", "Default user", [], True)],
    capabilities=[],
    decisions=[Decision("DEC-01", "Auth", "CAP-01", "", ["Pass"])],
    states=[State("S-00", "Start", "Start", ["DEC-01"], False)],
    tools=[Tool("Filing API", "CAP-01", True)],
)


def _pack(scenarios):
    path = os.path.join(tempfile.mkdtemp(), "pack.xlsx")
    write_challenge_pack(path, _INTAKE, scenarios)
    return load_workbook(path)


class TestChallengePack(unittest.TestCase):
    def setUp(self):
        self.scenarios = build_probes(_INTAKE)
        self.workbook = _pack(self.scenarios)

    def test_sheet_set_is_the_agreed_contract(self):
        self.assertEqual(self.workbook.sheetnames,
                         ["Instructions", "Scenarios", "Turn_Plan", "Run_Log", "Run_Summary"])

    def test_no_ground_truth_columns_anywhere(self):
        banned = {"decision path", "expected variant", "expected tool call", "category",
                  "materiality", "origin", "probe id", "probe family", "next state"}
        for name in self.workbook.sheetnames:
            headers = {str(c.value).strip().lower()
                       for c in next(self.workbook[name].iter_rows(max_row=1)) if c.value}
            self.assertFalse(headers & banned, f"{name} exposes {headers & banned}")

    def test_probe_expectations_never_appear_in_the_pack(self):
        """The behavioural assertion is the answer key — it must stay MRMG-internal."""
        blob = " ".join(str(cell) for name in self.workbook.sheetnames
                        for row in self.workbook[name].iter_rows(values_only=True)
                        for cell in row if cell)
        for scenario in self.scenarios:
            self.assertNotIn(scenario.termination, blob)
            self.assertNotIn(scenario.probe_id, blob)
            self.assertNotIn(scenario.probe_family, blob)

    def test_run_log_is_prepopulated_to_runs_times_turns(self):
        expected = sum(required_runs(s.effective_materiality) * recommended_turns(s)
                       for s in self.scenarios)
        self.assertEqual(self.workbook["Run_Log"].max_row - 1, expected)

    def test_run_summary_is_prepopulated_to_one_row_per_run(self):
        expected = sum(required_runs(s.effective_materiality) for s in self.scenarios)
        self.assertEqual(self.workbook["Run_Summary"].max_row - 1, expected)

    def test_review_verdict_drives_the_run_count(self):
        """The pack must reflect the review, not the tier the first sweep assigned."""
        scenario = build_probes(_INTAKE)[0]
        scenario.materiality, scenario.review_materiality = "Low", "Critical"
        rows = _pack([scenario])["Run_Summary"].max_row - 1
        self.assertEqual(rows, required_runs("Critical"))

    def test_proposed_scenarios_reach_the_pack_without_their_provenance(self):
        scenario = build_probes(_INTAKE)[0]
        scenario.origin = "llm-proposed"
        scenario.proposed_rationale = "Fills a gap in coverage."
        scenario.proposed_anchor = "SC-004"
        workbook = _pack([scenario])
        blob = " ".join(str(cell) for name in workbook.sheetnames
                        for row in workbook[name].iter_rows(values_only=True)
                        for cell in row if cell)
        self.assertNotIn("llm-proposed", blob)
        self.assertNotIn("Fills a gap in coverage.", blob)

    def test_materiality_override_drives_the_run_count(self):
        scenario = build_probes(_INTAKE)[0]
        scenario.materiality, scenario.materiality_override = "Low", "Critical"
        rows = _pack([scenario])["Run_Summary"].max_row - 1
        self.assertEqual(rows, required_runs("Critical"))


if __name__ == "__main__":
    unittest.main()
