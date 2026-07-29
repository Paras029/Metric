"""Minimal spreadsheet helpers: a bold header row, column widths, plain data rows.

Unstyled beyond a bold header and a frozen top row, keeping templates and outputs
easy to read and to edit by hand.
"""
from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
from typing import Iterable, Iterator, List

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_BOLD = Font(bold=True)


@contextmanager
def open_for_reading(path: str, description: str = "workbook") -> Iterator:
    """A workbook opened to be read and nothing else, and reliably closed afterwards.

    ``read_only`` puts openpyxl into its streaming reader rather than building a cell object for
    every cell in the file up front, which is most of the cost of opening a large registry -- a
    three-hundred-scenario one takes about a fifth as long this way, and that read happens on
    every page load of every scenario stage.

    It also holds the file open until told otherwise, which matters more here than the speed
    does: this tool is routinely run against a OneDrive-synced folder on Windows, where a handle
    left open is what turns the next write into a permission error. Hence a context manager
    rather than a bare loader -- there is no path out of this, including an exception mid-read,
    that leaves the file held.
    """
    try:
        workbook = load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(
            f"'{Path(path).name}' could not be opened as {description} ({exc}). A .xlsx file is a "
            f"zip archive underneath; anything else carrying that extension -- an older .xls saved "
            f"with the wrong name, a download that did not finish, a password-protected file -- "
            f"fails here. Re-save an unprotected copy from Excel (File > Save As > Excel "
            f"Workbook).") from exc
    try:
        yield workbook
    finally:
        workbook.close()


def add_sheet(workbook, name: str, headers: List[str], widths: List[int] = None) -> Worksheet:
    """Create a sheet with a bold, frozen header row."""
    sheet = workbook.create_sheet(name)
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = _BOLD
    for column, width in enumerate(widths or [], start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    return sheet


def write_rows(sheet: Worksheet, rows: Iterable[Iterable], start: int = 2) -> None:
    """Write data rows starting just below the header."""
    for r, row in enumerate(rows, start=start):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=r, column=column, value=value)


def read_rows(sheet: Worksheet) -> List[List[str]]:
    """Read every non-empty data row (row 1 is the header) as trimmed strings."""
    return read_table(sheet)[1]


def read_table(sheet: Worksheet) -> "tuple[List[str], List[List[str]]]":
    """The header row and the data rows, in one pass over the sheet.

    One pass because a read-only workbook streams: asking for the header and then the rows would
    mean opening the sheet twice. Returning the header at all is what lets a reader address
    columns by the name printed in the file rather than by counting them -- see
    :mod:`scenario_generator.io.workbooks`, where a registry has twenty-six of them and hand-
    counted indices had to be kept in step with the writer by eye.
    """
    header: List[str] = []
    rows: List[List[str]] = []
    for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        if number == 1:
            header = ["" if v is None else str(v).strip() for v in (values or ())]
            continue
        if values and values[0] not in (None, ""):
            rows.append(["" if v is None else str(v).strip() for v in values])
    return header, rows
