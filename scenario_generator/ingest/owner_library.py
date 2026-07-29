"""Reading the model owner's own test scenarios out of whatever they sent.

Nobody submits their testing in the shape this tool would like. It arrives as a spreadsheet with
the headings someone chose two years ago, a Word table, a CSV export, or a numbered list in a
document. Refusing anything but one exact layout would put the coverage measurement out of reach
for most submissions, which is the same as not having it.

So the reader works by recognising rather than requiring: it finds the sheet that looks like a
scenario list, works out which column holds the identifier and which holds the description from
the headings and the shape of the data, and falls back to reading numbered paragraphs where there
is no table at all.

What it will not do is guess quietly. Every file records how it was read, and a file it could not
make sense of is reported rather than returned empty -- an owner who appears to have tested
nothing, when in fact their format was not understood, is a conclusion worth avoiding.
"""
from __future__ import annotations

import csv
import logging
import re
from pathlib import Path
from typing import List, Tuple

from ..core.models import OwnerScenario

logger = logging.getLogger(__name__)

# Headings that name the identifier, the description, and the route, in the words teams actually
# use. Matched case-insensitively as substrings, longest first.
_ID_HEADINGS = ("scenario id", "test id", "case id", "test case id", "tc id", "ref", "id", "no",
                "number", "s.no", "sr no")
_DESCRIPTION_HEADINGS = ("scenario description", "test scenario", "description", "scenario",
                         "test case", "summary", "objective", "narrative", "steps", "title",
                         "name")
_PATH_HEADINGS = ("decision path", "path", "flow", "expected route", "route")

# A line that opens with its own number, for documents with no table at all.
_NUMBERED = re.compile(r"^\s*(?:(\d+)[.)]|[-*•])\s+(.{15,})$")

MIN_DESCRIPTION_CHARS = 15


class UnreadableLibrary(Exception):
    """The submitted file could not be read as a list of scenarios, with a reason."""


def read_owner_library(path: Path) -> Tuple[List[OwnerScenario], str]:
    """Read a submitted scenario library. Returns the scenarios and how they were found."""
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in (".xlsx", ".xlsm"):
        return _from_workbook(path)
    if suffix == ".csv":
        return _from_csv(path)
    if suffix == ".docx":
        return _from_docx(path)
    if suffix in (".md", ".txt"):
        return _from_lines(path.read_text(encoding="utf-8", errors="replace").splitlines(),
                           f"numbered lines in {path.name}")
    if suffix == ".pdf":
        return _from_pdf(path)

    raise UnreadableLibrary(
        f"{suffix or 'this file type'} cannot be read as a scenario library. Supported: "
        f".xlsx, .xlsm, .csv, .docx, .pdf, .md, .txt.")


# --------------------------------------------------------------------------- tables
def _score_header(cells: List[str]) -> int:
    """How much a row looks like a header rather than data."""
    text = " ".join(cells).lower()
    return sum(1 for word in ("id", "description", "scenario", "test", "expected", "step")
               if word in text)


def _columns(header: List[str]) -> Tuple[int, int, int]:
    """Which columns hold the id, the description and the route.

    Longest heading first, so "scenario description" is not claimed by "scenario", and the
    description falls back to the widest column when no heading matches -- a scenario list
    without a recognisable heading still has one column carrying far more text than the rest.
    """
    lowered = [c.strip().lower() for c in header]

    def find(candidates) -> int:
        for candidate in sorted(candidates, key=len, reverse=True):
            for index, cell in enumerate(lowered):
                if candidate == cell or candidate in cell:
                    return index
        return -1

    return find(_ID_HEADINGS), find(_DESCRIPTION_HEADINGS), find(_PATH_HEADINGS)


def _from_rows(rows: List[List[str]], origin: str) -> Tuple[List[OwnerScenario], str]:
    """Turn a table into scenarios, working out its shape from the first rows."""
    rows = [[str(cell or "").strip() for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        raise UnreadableLibrary(f"{origin} has no rows.")

    # The header is rarely the first row -- spreadsheets carry titles and blank spacers.
    header_index = max(range(min(5, len(rows))), key=lambda i: _score_header(rows[i]))
    id_col, description_col, path_col = _columns(rows[header_index])
    body = rows[header_index + 1:]

    if description_col < 0:
        widths = [sum(len(row[i]) for row in body if i < len(row))
                  for i in range(max((len(r) for r in body), default=0))]
        if not widths:
            raise UnreadableLibrary(f"{origin} has a header but no data rows.")
        description_col = widths.index(max(widths))
        how = f"{origin}, widest column taken as the description"
    else:
        how = f"{origin}, column '{rows[header_index][description_col]}'"

    scenarios = []
    for number, row in enumerate(body, start=1):
        description = row[description_col] if description_col < len(row) else ""
        if len(description) < MIN_DESCRIPTION_CHARS:
            continue                                       # a heading, a spacer, or a stray cell
        identifier = (row[id_col].strip() if 0 <= id_col < len(row) and row[id_col].strip()
                      else f"OS-{number:03d}")
        declared = row[path_col] if 0 <= path_col < len(row) else ""
        scenarios.append(OwnerScenario(identifier, description, declared))

    if not scenarios:
        raise UnreadableLibrary(
            f"{origin} was read but no row held a description of at least "
            f"{MIN_DESCRIPTION_CHARS} characters.")
    return scenarios, how


def _from_workbook(path: Path) -> Tuple[List[OwnerScenario], str]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(str(path), data_only=True)
    except Exception as exc:
        # A .xlsx file is a zip archive; anything else under that extension fails here with a
        # message that names the container format rather than the fix. The likely causes are all
        # ones the sender can act on: an older .xls saved under the wrong extension, a download
        # that did not finish, or a password-protected file -- openpyxl cannot open any of those,
        # and the raw error ("File is not a zip file") does not say so.
        raise UnreadableLibrary(
            f"{path.name} could not be opened as an Excel workbook ({exc}). This usually means "
            f"the file is not really .xlsx underneath -- an older .xls saved with the wrong "
            f"extension, a download that did not finish, or a password-protected file. Re-save an "
            f"unprotected copy from Excel (File > Save As > Excel Workbook), or send it as .csv, "
            f"which this reads just as well.") from exc

    # Try each sheet and keep whichever yields the most scenarios. A workbook usually carries a
    # cover sheet, a glossary and the actual list, and the list is the longest.
    best: Tuple[List[OwnerScenario], str] = ([], "")
    problems = []
    for name in book.sheetnames:
        rows = [[cell for cell in row] for row in book[name].iter_rows(values_only=True)]
        try:
            found, how = _from_rows(rows, f"sheet '{name}'")
        except UnreadableLibrary as exc:
            problems.append(str(exc))
            continue
        if len(found) > len(best[0]):
            best = (found, how)

    if not best[0]:
        raise UnreadableLibrary(
            "no sheet in the workbook looked like a scenario list. " + " ".join(problems[:2]))
    return best


def _from_csv(path: Path) -> Tuple[List[OwnerScenario], str]:
    with path.open(newline="", encoding="utf-8", errors="replace") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [list(row) for row in csv.reader(handle, dialect)]
    return _from_rows(rows, path.name)


def _from_docx(path: Path) -> Tuple[List[OwnerScenario], str]:
    from docx import Document

    document = Document(str(path))

    # A table is the more structured statement of the same thing, so prefer it.
    for index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        try:
            return _from_rows(rows, f"table {index} in {path.name}")
        except UnreadableLibrary:
            continue

    return _from_lines([p.text for p in document.paragraphs],
                       f"numbered paragraphs in {path.name}")


def _from_pdf(path: Path) -> Tuple[List[OwnerScenario], str]:
    from pypdf import PdfReader

    lines: List[str] = []
    for page in PdfReader(str(path)).pages:
        lines.extend((page.extract_text() or "").splitlines())
    return _from_lines(lines, f"numbered lines in {path.name}")


def _from_lines(lines: List[str], origin: str) -> Tuple[List[OwnerScenario], str]:
    """Read a numbered or bulleted list, for submissions with no table at all."""
    scenarios = []
    for line in lines:
        match = _NUMBERED.match(line or "")
        if not match:
            continue
        number, description = match.group(1), match.group(2).strip()
        if len(description) < MIN_DESCRIPTION_CHARS:
            continue
        identifier = f"OS-{int(number):03d}" if number else f"OS-{len(scenarios) + 1:03d}"
        scenarios.append(OwnerScenario(identifier, description, ""))

    if not scenarios:
        raise UnreadableLibrary(
            f"{origin} contained no numbered or bulleted lines that read as scenarios.")
    return scenarios, origin
