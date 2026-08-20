"""Read a filled intake workbook, read an owner's declared scenario list, and write the blank
template. Assumes a well-formed workbook, with row 1 of each sheet as the header.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook, load_workbook

from metric.domain import sheets
from metric.shared.text import ID_BODY, is_yes, normalise_variant, one_of, parse_reached_via, split_list
from metric.domain.models import CATEGORIES, INPUT_SOURCES, Capability, Decision, IntakeData, OwnerScenario, Persona, State, Tool

_DECISION_TOKEN = re.compile(r"DEC-" + ID_BODY)
_STATE_TOKEN = re.compile(r"S-" + ID_BODY, re.I)

# The sheets an intake must have. "Tools" is optional -- an agent that calls nothing
# is unusual but not malformed.
_REQUIRED_SHEETS = ("L1 Use Case", "Personas", "L2 Capabilities", "L3 Decisions",
                    "L4 States")


def _open_for_editing(path: str):
    """Open a workbook that is about to be written back to."""
    try:
        return load_workbook(path)
    except Exception as exc:
        raise ValueError(
            f"'{Path(path).name}' could not be opened as an Excel workbook ({exc}). This "
            f"usually means the file is not really .xlsx underneath -- an older .xls saved with "
            f"the wrong extension, a download that did not finish, or a password-protected file. "
            f"Re-save an unprotected copy from Excel (File > Save As > Excel Workbook), then "
            f"upload that.") from exc


def _cell(row: List[str], index: int) -> str:
    return row[index] if index < len(row) else ""


def _named(cell: str, declared: List[str], shape) -> tuple:
    """The ids a cell names, in the order they are written."""
    text = str(cell or "")
    if not text.strip():
        return ()

    canonical = {i.upper(): i for i in declared if i}
    # Longest first, or "S-1" would match the front of "S-10" and leave the rest behind.
    known = sorted(canonical.values(), key=len, reverse=True)
    expression = re.compile(
        r"(?<![A-Za-z0-9_-])(?:" + "|".join([re.escape(i) for i in known] + [shape.pattern])
        + r")(?![A-Za-z0-9_-])", re.I)

    found = []
    for token in expression.findall(text):
        settled = canonical.get(token.upper(), token.upper())
        if settled not in found:
            found.append(settled)
    return tuple(found)


def _input_source(raw: str) -> str:
    """Match the declared source against the closed vocabulary; anything unknown reads as User."""
    return one_of(raw, INPUT_SOURCES, "User")


def _max_attempts(raw: str) -> int:
    """How many times one decision may be exercised within a single path. Blank or invalid -> 1."""
    try:
        return max(1, int(float(str(raw).strip())))
    except (TypeError, ValueError):
        return 1


def _outcome_type(raw: str) -> str:
    """Declared category of a terminal state; blank falls back to the keyword heuristic."""
    return one_of(raw, CATEGORIES)


# --------------------------------------------------------------------------- readers
def _sheet(workbook, name: str, path: str) -> list:
    """One sheet's data rows, or a message naming the sheet that is not there."""
    if name not in workbook.sheetnames:
        raise ValueError(
            f"'{Path(path).name}' has no '{name}' sheet, so it is not an intake workbook this "
            f"can read. It needs the sheets the template ships with: "
            f"{', '.join(_REQUIRED_SHEETS)}. Download a blank template and fill that in, or "
            f"check the right file was uploaded.")
    return sheets.read_rows(workbook[name])


def read_intake(path: str) -> IntakeData:
    with sheets.open_for_reading(path, "an intake workbook") as workbook:
        rows = {name: _sheet(workbook, name, path) for name in _REQUIRED_SHEETS}
        rows["Tools"] = (sheets.read_rows(workbook["Tools"])
                         if "Tools" in workbook.sheetnames else [])

    use_case = {_cell(r, 0): _cell(r, 1) for r in rows["L1 Use Case"] if _cell(r, 0)}

    personas = [Persona(_cell(r, 0), _cell(r, 1),
                        split_list(_cell(r, 2), separators=r"[,;]"), is_yes(_cell(r, 3)))
                for r in rows["Personas"]]
    if not personas:
        raise ValueError(
            f"'{Path(path).name}' declares no personas, and every scenario is walked by one. Add "
            f"at least one row to the Personas sheet -- a cooperative user trying to use the "
            f"service as intended is the one every agent has.")
    if not any(p.is_default for p in personas):
        personas[0] = Persona(personas[0].id, personas[0].name, personas[0].applies_to, True)

    decisions = [Decision(_cell(r, 0), _cell(r, 1), _cell(r, 2), _cell(r, 3),
                          [normalise_variant(v) for v in split_list(_cell(r, 4), separators=r"[/]")],
                          _input_source(_cell(r, 5)), _max_attempts(_cell(r, 6)), _cell(r, 7),
                          is_yes(_cell(r, 8)))
                 for r in rows["L3 Decisions"]]

    states = [State(_cell(r, 0), _cell(r, 1), _cell(r, 2),
                    list(_named(_cell(r, 3), [d.id for d in decisions], _DECISION_TOKEN)),
                    is_yes(_cell(r, 4)), _outcome_type(_cell(r, 5)))
              for r in rows["L4 States"]]

    # Read after the states, so a span can be matched against the ids that actually exist.
    #
    # Entry and exit are read as whatever state ids appear in the cell, so "S-00, S-03" and
    # "S-00 and S-03" and a list down the cell all mean the same thing. The alternative is a
    # separator nobody remembers and a span that silently covers half the block it names.
    declared = [s.id for s in states]
    capabilities = [Capability(_cell(r, 0), _cell(r, 1), _cell(r, 2),
                               _named(_cell(r, 3), declared, _STATE_TOKEN),
                               _named(_cell(r, 4), declared, _STATE_TOKEN))
                    for r in rows["L2 Capabilities"]]

    tools = [Tool(_cell(r, 0), _cell(r, 1), is_yes(_cell(r, 2)))
             for r in rows["Tools"] if _cell(r, 0)]

    return IntakeData(use_case, personas, capabilities, decisions, states, tools)


def read_review_notes(path: str) -> List[dict]:
    """The drafter's own review notes, read back from the "Review This" sheet it wrote."""
    try:
        with sheets.open_for_reading(path, "an intake workbook") as workbook:
            if "Review This" not in workbook.sheetnames:
                return []
            rows = sheets.read_rows(workbook["Review This"])
    except Exception:
        return []

    notes: List[dict] = []
    in_notes = False
    for row in rows:
        first = _cell(row, 0).strip()
        if first == "Note":
            in_notes = True
            continue
        if not in_notes:
            continue
        note = _cell(row, 2).strip()
        if first and note:
            notes.append({"field": first, "note": note})
    return notes


def read_owner_scenarios(path: str, sheet_name: str = "Scenarios") -> List[OwnerScenario]:
    """Read a model owner's declared scenario list: ID, Description, optional Decision Path."""
    with sheets.open_for_reading(path, "a scenario library") as workbook:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"'{Path(path).name}' has no '{sheet_name}' sheet. This reads a list with one row "
                f"per scenario: an id in the first column, a description in the second, and "
                f"optionally a decision path in the third. It has: "
                f"{', '.join(workbook.sheetnames)}.")
        rows = sheets.read_rows(workbook[sheet_name])
    return [OwnerScenario(_cell(r, 0), _cell(r, 1), _cell(r, 2))
            for r in rows if _cell(r, 0) and _cell(r, 1)]


# --------------------------------------------------------------------------- template
_GUIDE = [
    ("Fill order", "L1 Use Case, Personas, L2 Capabilities, L3 Decisions, L4 States, then Tools."),
    ("The core idea", "L3 Decisions and L4 States describe the agent as a graph. A State is a "
                      "position the conversation can be in; a Decision is a branch point with named "
                      "outcomes; 'Reached Via' on a State says which Decision outcome leads there."),
    ("Entry / Exit States", "On L2 Capabilities: which states a capability is entered in, and "
                            "which it hands on or finishes in. Scenarios are enumerated one "
                            "capability at a time, from each entry state to any exit state, so an "
                            "exit of one capability is usually an entry of the next. Leave both "
                            "blank on every capability to walk the whole graph end to end instead."),
    ("What connects the graph", "Two columns, and only two. 'Reached Via' on a State says which "
                                "decision outcome arrives there; 'Valid Next Decisions' says "
                                "which decisions leave it. Every route the tool ever walks runs "
                                "along those. A decision no state lists under Valid Next "
                                "Decisions is never reached; a state no outcome names in Reached "
                                "Via is never arrived at. Both are drawn as fragments below the "
                                "graph rather than in it, and neither contributes a scenario."),
    ("Reached Via", "Use 'Start' for the opening state, or 'DEC-xx=Variant' for a state reached by a "
                    "decision outcome -- the variant spelled exactly as that decision spells it. "
                    "Two states may share the same value if an outcome recurs, and one state may "
                    "name several outcomes comma-separated where routes converge: "
                    "'DEC-02=Too old, DEC-05=Withdrawn'."),
    ("Valid Next Decisions", "Which decisions can be taken from this state, by id, "
                             "comma-separated. Empty only where the interaction ends here -- an "
                             "intermediate state with nothing next stops every route through it."),
    ("Possible Outputs", "Separate outcomes with '/'. Use plain words where they fit (Pass, Fail, "
                         "Escalate, Terminate, Retry) so scenarios are categorised correctly."),
    ("Input Source", "Where the decision's input arrives from: User, Tool, Memory-Session, "
                     "Memory-CrossSession, System-Context, Document. Only 'User' steps become "
                     "conversational turns — internal reasoning, memory lookups and tool-driven "
                     "branches are still decisions, but the tester does not script them. Blank "
                     "reads as User."),
    ("Max Attempts", "How many times this one decision may be exercised within a single path, "
                     "e.g. 3 where authentication allows three tries. Blank reads as 1. This is "
                     "what bounds retry loops — there is no global cap."),
    ("Outcome Condition", "Optional. What actually triggers each outcome, e.g. 'risk score < 0.3'. "
                          "Not used for graph traversal; recorded so boundary inputs can be "
                          "derived later."),
    ("Out of Scope?", "Y if this decision has already been reviewed elsewhere and is being reused "
                      "as-is -- a plug-and-play sub-system covered by a separate engagement is "
                      "the usual case. The decision still belongs in the sheet, because the graph "
                      "is not honest without it, but no scenario is generated through it. Blank "
                      "reads as No."),
    ("Outcome Type", "On a terminal state only: Happy path, Retry, Fallback, Escalation or "
                     "Termination. This is what sets a scenario's category, so declare it rather "
                     "than relying on outcome wording."),
    ("Capability Type", "One of: Lookup (reads and reports), Transactional (changes account or "
                        "financial state), Gating (authenticates or authorises), Advisory "
                        "(explains or recommends), PII-handling (touches sensitive personal "
                        "data). Used to decide which adversarial probes apply, so pick the "
                        "closest fit rather than leaving it blank."),
    ("Example — Persona", "P1 | Cooperative, verified user | Happy path | Y"),
    ("Example — Capability", "CAP-01 | Authentication | Gating"),
    ("Example — Decision", "DEC-01 | Authentication outcome | CAP-01 | credentials | Pass / Fail "
                           "| User | 3 | 3 failed tries locks the session"),
    ("Example — State", "S-00 | Start | Session begins, unauthenticated | DEC-01 | No | (blank)"),
    ("Example — Tool", "Identity verification service | CAP-01 | No"),
]

_L1_FIELDS = [
    "Use case name", "Business objective", "Use case rating (informational)",
    "Agent type", "Channel / modality", "Human handoff triggers",
    "Safety requirements", "Success criteria", "Known limitations",
]


_NEXT_DECISIONS_COLUMN = 4                                 # "Valid Next Decisions" on L4 States


def set_decision_scope(path: str, decision_id: str, out_of_scope: bool) -> bool:
    """Flip one decision's Out of Scope column in place. Returns whether a row was found."""
    workbook = _open_for_editing(path)
    if "L3 Decisions" not in workbook.sheetnames:
        return False
    sheet = workbook["L3 Decisions"]
    for row in sheet.iter_rows(min_row=2):
        if row and str(row[0].value or "").strip() == decision_id:
            sheet.cell(row=row[0].row, column=9, value="Yes" if out_of_scope else "No")
            workbook.save(path)
            return True
    return False


def set_capability_span(path: str, capability_id: str, entry_states: List[str],
                        exit_states: List[str]) -> bool:
    """Set one capability's entry and exit states in place. Returns whether a row was found."""
    workbook = _open_for_editing(path)
    if "L2 Capabilities" not in workbook.sheetnames:
        return False
    sheet = workbook["L2 Capabilities"]
    for row in sheet.iter_rows(min_row=2):
        if row and str(row[0].value or "").strip() == capability_id:
            sheet.cell(row=row[0].row, column=4, value=", ".join(entry_states))
            sheet.cell(row=row[0].row, column=5, value=", ".join(exit_states))
            workbook.save(path)
            return True
    return False


def set_state_reached_via(path: str, state_id: str, reached_via: str) -> bool:
    """Correct one state's Reached Via column in place. Returns whether a row was found."""
    workbook = _open_for_editing(path)
    if "L4 States" not in workbook.sheetnames:
        return False
    sheet = workbook["L4 States"]
    for row in sheet.iter_rows(min_row=2):
        if row and str(row[0].value or "").strip() == state_id:
            sheet.cell(row=row[0].row, column=2, value=reached_via)
            workbook.save(path)
            return True
    return False


def attach_decision_to_state(path: str, state_id: str, decision_id: str) -> bool:
    """Add ``decision_id`` to a state's Valid Next Decisions. Returns whether anything changed."""
    workbook = _open_for_editing(path)
    if "L4 States" not in workbook.sheetnames:
        return False

    sheet = workbook["L4 States"]
    for row in sheet.iter_rows(min_row=2):
        if str(row[0].value or "").strip() != state_id:
            continue
        cell = row[_NEXT_DECISIONS_COLUMN - 1]
        already = _DECISION_TOKEN.findall(str(cell.value or ""))
        if decision_id in already:
            return False
        cell.value = ", ".join(already + [decision_id])
        workbook.save(path)
        return True
    return False


def write_template(path: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    guide = sheets.add_sheet(workbook, "Guide", ["Topic", "Notes"], [26, 100])
    sheets.write_rows(guide, _GUIDE)

    l1 = sheets.add_sheet(workbook, "L1 Use Case", ["Field", "Value"], [34, 90])
    sheets.write_rows(l1, [[field, ""] for field in _L1_FIELDS])

    sheets.add_sheet(workbook, "Personas",
                     ["ID", "Name", "Applies To", "Default"], [10, 34, 28, 10])
    sheets.add_sheet(workbook, "L2 Capabilities",
                     ["Capability ID", "Name", "Type", "Entry States", "Exit States"],
                     [14, 30, 20, 22, 26])
    sheets.add_sheet(workbook, "L3 Decisions",
                     ["Decision ID", "Decision", "Triggering Capability", "Inputs",
                      "Possible Outputs", "Input Source", "Max Attempts", "Outcome Condition",
                      "Out of Scope?"],
                     [12, 28, 22, 30, 30, 20, 14, 34, 14])
    sheets.add_sheet(workbook, "L4 States",
                     ["State ID", "Reached Via", "Description", "Valid Next Decisions",
                      "Terminal?", "Outcome Type"], [10, 26, 34, 26, 11, 16])
    sheets.add_sheet(workbook, "Tools",
                     ["Tool Name", "Capability ID", "State-changing?"], [30, 16, 16])

    workbook.save(path)
